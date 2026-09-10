from pathlib import Path
from datetime import datetime,timedelta
import numpy as np,json
from openpyxl import load_workbook
root=Path('/Users/justingao/Documents/CUMCM/CUMCM2026Problems/C题/附件')
def read_matrix(file,sheet):
 w=load_workbook(root/file,read_only=True,data_only=True);s=w[sheet]
 rows=list(s.iter_rows(values_only=True));w.close()
 return rows[0],rows[1:],np.array([r[1:] for r in rows[1:]],float)
_,lr,L=read_matrix('附件2.xlsx','小区负载')
_,_,G=read_matrix('附件2.xlsx','光伏发电实际功率')
_,_,P=read_matrix('附件4.xlsx','Sheet1')
def summary(a):
 m=a[:31];v=m.reshape(-1)
 pattern=np.broadcast_to(m.mean(0),m.shape)
 return {'january_count':int(v.size),'january_day_lag_corr':float(np.corrcoef(v[144:],v[:-144])[0,1]),'january_week_lag_corr':float(np.corrcoef(v[1008:],v[:-1008])[0,1]),'january_time_of_day_in_sample_variance_fraction':float(1-((m-pattern)**2).sum()/((m-m.mean())**2).sum()),'january_min':float(v.min()),'january_max':float(v.max()),'january_zero_fraction':float((v==0).mean())}
out={'scope':'January only for structural diagnostics; no prediction training or optimization. Time-of-day fit is in-sample description, not predictive accuracy.','load':summary(L),'pv':summary(G),'price':summary(P)}
actual={}
for i,row in enumerate(lr):
 day=row[0]
 for t in range(144):actual[day+timedelta(minutes=10*(t+1))]=G[i,t]
w=load_workbook(root/'附件3.xlsx',read_only=True,data_only=True)
pred={};date=None
for row in w.active.iter_rows(min_row=2,values_only=True):
 if row[0]:
  if isinstance(row[0],datetime):date=row[0]
  else:date=datetime.strptime(str(row[0]),'%Y-%m-%d')
 hour=int(str(row[1]).split(':')[0]);issue=date+timedelta(hours=hour)
 if date.month!=1:continue
 for h,val in enumerate(row[2:],1):pred[(date,hour,issue+timedelta(hours=h))]=float(val)
w.close()
res={str(h):[] for h in [0,6,12]}
for d in range(31):
 day=datetime(2025,1,1)+timedelta(days=d)
 for hr in range(13,19):
  target=day+timedelta(hours=hr)
  if target not in actual:continue
  for issuehour in [0,6,12]:res[str(issuehour)].append(pred[(day,issuehour,target)]-actual[target])
out['pv_paired_forecast_january_13_to_18_hour_labels']={k:{'n':len(v),'mae_kw':float(np.mean(np.abs(v))),'bias_forecast_minus_actual_kw':float(np.mean(v))} for k,v in res.items()}
out['forecast_diagnostic_caveat']='Matches nominal whole-hour labels only; does not resolve whether actual cells are point values or interval averages. Same target hours/dates for three releases. Descriptive, not proof of economic benefit.'
p=Path('/Users/justingao/Documents/CUMCM/选题分析/C-模型选择数据证据.json');p.write_text(json.dumps(out,ensure_ascii=False,indent=2));print(json.dumps(out,ensure_ascii=False,indent=2))

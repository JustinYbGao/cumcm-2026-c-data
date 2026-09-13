"""Deterministic parsing and explicitly assumed energy conversion; no modeling."""
import argparse
from datetime import datetime, time
import json
from pathlib import Path
import re
import numpy as np
import pandas as pd
import openpyxl
from openpyxl.utils import get_column_letter

WORK = Path(__file__).resolve().parents[1]
ROOT = WORK
CONFIG = json.loads((WORK / 'configs/data_processing.json').read_text())
SOURCE = ROOT / CONFIG['source_directory']
OUT = WORK / 'data/processed'


def minute_of_day(value):
    if isinstance(value, (time, datetime, pd.Timestamp)):
        if value.second or value.microsecond:
            raise ValueError(f'Non-minute time: {value}')
        return value.hour * 60 + value.minute
    match = re.fullmatch(r'(\d{1,2}):(\d{1,2})(?::00)?(\+1)?', str(value).strip())
    if not match:
        raise ValueError(f'Invalid time label: {value!r}')
    h, m = int(match[1]), int(match[2])
    if m >= 60 or h > 24 or (h == 24 and m):
        raise ValueError(f'Invalid clock time: {value}')
    return h * 60 + m + (1440 if match[3] else 0)


def naive_time(value):
    t = pd.Timestamp(value)
    if pd.isna(t) or t.tzinfo is not None:
        raise ValueError('Decision time must be a nonempty timezone-naive problem-local timestamp')
    return t


def history_at(actual, decision_time):
    """Completed records, including the interval ending exactly at decision_time."""
    t = naive_time(decision_time)
    return actual.loc[pd.to_datetime(actual.interval_end) <= t].copy()


def forecasts_at(forecast, decision_time):
    """All available versions for future targets; no implicit latest-version collapse."""
    t = naive_time(decision_time)
    future = (pd.to_datetime(forecast.interval_start) >= t if 'interval_start' in forecast
              else pd.to_datetime(forecast.target_time) > t)
    return forecast.loc[(pd.to_datetime(forecast.issue_time) <= t) & future].copy()


def rows_of(filename, sheet):
    wb = openpyxl.load_workbook(SOURCE / filename, read_only=True, data_only=True)
    rows = list(wb[sheet].values)
    wb.close()
    return rows


def numeric(values):
    # Refuse unexpected blanks/text; zero and every finite signed value are preserved.
    array = np.asarray(values)
    if any(not isinstance(v, (int, float)) or isinstance(v, bool) for v in array.flat):
        raise ValueError('Blank or nonnumeric source value; inspect original before changing it')
    array = array.astype(float)
    if not np.isfinite(array).all():
        raise ValueError('Nonfinite source value')
    return array


def slots(labels):
    minutes = [minute_of_day(v) for v in labels]
    if minutes != list(range(10, 1441, 10)):
        raise ValueError('Source time labels do not cover ordered 00:10 through 24:00')
    return np.asarray(minutes)


def write(frame, name):
    frame.to_csv(OUT / f'{name}.csv', index=False, encoding='utf-8', date_format='%Y-%m-%dT%H:%M:%S')
    print(f'Wrote {name}: {len(frame)} rows', flush=True)


def prepare_q1():
    rows = rows_of('attachment_1.xlsx', 'Sheet1')
    if rows[0] != ('\u65f6\u95f4', '\u7535\u4ef7', '\u5c0f\u533a\u8d1f\u8f7d', '\u5149\u4f0f\u53d1\u7535\u9884\u6d4b\u529f\u7387'):
        raise ValueError('Unexpected attachment 1 header')
    labels = [r[0] for r in rows[1:]]
    minutes = slots(labels)
    q = pd.DataFrame({'slot_id': np.arange(1, len(minutes)+1),
                      'raw_time_label': [str(v) for v in labels],
                      'start_minute': minutes-10, 'end_minute': minutes})
    # Q1 has no source date. Do not attach a fictional full timestamp to this day.
    for field, mins in [('start_time', minutes-10), ('end_time', minutes)]:
        q[field] = [f'{m//60:02d}:{m%60:02d}:00' for m in mins]
    q[['price_yuan_per_kwh', 'load_kw', 'pv_forecast_kw']] = numeric([r[1:] for r in rows[1:]])
    q['load_kwh'] = q.load_kw/6
    q['pv_forecast_kwh'] = q.pv_forecast_kw/6
    q['net_load_forecast_kw'] = q.load_kw-q.pv_forecast_kw
    q['net_load_forecast_kwh'] = q.load_kwh-q.pv_forecast_kwh
    q['source_row'] = np.arange(2, len(q)+2)
    fixed = q[['slot_id', 'raw_time_label', 'start_minute', 'end_minute',
               'start_time', 'end_time', 'price_yuan_per_kwh', 'source_row']].copy()
    write(q, 'q1_day')
    write(fixed, 'fixed_price')
    return q, fixed


def annual_long(filename, sheet, value_name, prefix):
    rows = rows_of(filename, sheet)
    labels = rows[0][1:]
    minutes = slots(labels)
    dates = pd.DatetimeIndex([naive_time(r[0]) for r in rows[1:]])
    if not dates.equals(pd.date_range('2025-01-01', '2025-12-31')):
        raise ValueError(f'Unexpected source dates: {filename}/{sheet}')
    n = len(dates)
    frame = pd.DataFrame({'date': np.repeat(dates.to_numpy(), 144),
                          'slot_id': np.tile(np.arange(1,145), n),
                          value_name: numeric([r[1:] for r in rows[1:]]).reshape(-1)})
    frame['interval_end'] = frame.date + pd.to_timedelta(np.tile(minutes, n), unit='m')
    frame['interval_start'] = frame.interval_end-pd.Timedelta(minutes=10)
    frame[f'{prefix}_raw_time_label'] = np.tile([str(v) for v in labels], n)
    frame[f'{prefix}_source_row'] = np.repeat(np.arange(2,n+2), 144)
    frame[f'{prefix}_source_column'] = np.tile([get_column_letter(c) for c in range(2,146)], n)
    return frame


def prepare_actual(fixed):
    load = annual_long('attachment_2.xlsx', '\u5c0f\u533a\u8d1f\u8f7d', 'load_actual_kw', 'load')
    pv = annual_long('attachment_2.xlsx', '\u5149\u4f0f\u53d1\u7535\u5b9e\u9645\u529f\u7387', 'pv_actual_kw', 'pv')
    price = annual_long('attachment_4.xlsx', 'Sheet1', 'actual_price_yuan_per_kwh', 'price')
    keys = ['date','slot_id','interval_start','interval_end']
    actual = load.merge(pv, on=keys, how='outer', validate='one_to_one').merge(
        price, on=keys, how='outer', validate='one_to_one')
    actual = actual.merge(fixed[['slot_id','price_yuan_per_kwh']].rename(
        columns={'price_yuan_per_kwh':'fixed_price_yuan_per_kwh'}), on='slot_id', validate='many_to_one')
    if len(actual) != len(load) or actual.isna().any().any():
        raise ValueError('Unexpected merge expansion, loss or missing values')
    actual = actual.sort_values('interval_start').reset_index(drop=True)
    actual['load_actual_kwh'] = actual.load_actual_kw/6
    actual['pv_actual_kwh'] = actual.pv_actual_kw/6
    actual['net_load_actual_kw'] = actual.load_actual_kw-actual.pv_actual_kw
    actual['net_load_actual_kwh'] = actual.load_actual_kwh-actual.pv_actual_kwh
    actual['available_time'] = actual.interval_end
    actual['date'] = actual.date.dt.strftime('%Y-%m-%d')
    write(actual, 'actual_10min')
    return actual


def prepare_hourly():
    rows = rows_of('attachment_3.xlsx', 'Sheet1')
    if list(rows[0][2:]) != [f'\u9884\u62a5{h}\u5c0f\u65f6' for h in range(1,25)]:
        raise ValueError('Unexpected horizon headers')
    raw_dates = pd.Series([r[0] for r in rows[1:]], dtype=object)
    clean_dates = raw_dates.replace(r'^\s*$', np.nan, regex=True).ffill()
    dates = pd.to_datetime(clean_dates, errors='raise')
    minutes = [minute_of_day(r[1]) for r in rows[1:]]
    issues = pd.DatetimeIndex(dates+pd.to_timedelta(minutes, unit='m'))
    if not issues.equals(pd.date_range('2025-01-01', periods=1460, freq='6h')):
        raise ValueError('Unexpected issue dates/order')
    records = pd.DataFrame({'issue_time': np.repeat(issues.to_numpy(),24),
                            'horizon_h': np.tile(np.arange(1,25),len(issues)),
                            'pv_forecast_kw': numeric([r[2:] for r in rows[1:]]).reshape(-1)})
    records['target_time'] = records.issue_time+pd.to_timedelta(records.horizon_h,unit='h')
    records['source_row'] = np.repeat(np.arange(2,len(issues)+2),24)
    records['source_column'] = np.tile([get_column_letter(c) for c in range(3,27)],len(issues))
    records['raw_date_label'] = np.repeat(raw_dates.fillna('').to_numpy(),24)
    records['raw_issue_label'] = np.repeat([str(r[1]) for r in rows[1:]],24)
    records['date_was_filled'] = np.repeat((raw_dates.isna() | raw_dates.astype(str).str.strip().eq('')).to_numpy(),24)
    write(records, 'pv_forecast_hourly')
    changes = [{'source_file':'attachment_3.xlsx','source_sheet':'Sheet1','source_cell':f'A{i+2}',
                'original_value':json.dumps(raw_dates.iloc[i],ensure_ascii=False),
                'new_value':clean_dates.iloc[i], 'reason':'Forward-filled blank date labels within the existing issue order; no power value changed'}
               for i in range(len(raw_dates)) if records.date_was_filled.iloc[i*24]]
    pd.DataFrame(changes).to_csv(WORK/'data/interim/transformation_log.csv',index=False)
    return records


def convert_forecasts(hourly):
    """Point-power linear integral; first-hour anchor comes only from an older forecast."""
    frames, available_points = [], {}
    for issue, group in hourly.sort_values(['issue_time','horizon_h']).groupby('issue_time',sort=True):
        powers = group.pv_forecast_kw.to_numpy()
        anchor = available_points.get(issue)
        if anchor is None:
            knots = np.arange(1,25)*60
            p = powers
            starts = np.arange(60,1440,10)
        else:
            knots = np.arange(25)*60
            p = np.r_[anchor[1],powers]
            starts = np.arange(0,1440,10)
        left = np.interp(starts,knots,p)
        right = np.interp(starts+10,knots,p)
        out = pd.DataFrame({'issue_time':issue,
            'interval_start':issue+pd.to_timedelta(starts,unit='m'),
            'interval_end':issue+pd.to_timedelta(starts+10,unit='m'),
            'horizon_start_min':starts,'horizon_end_min':starts+10,
            'pv_forecast_kwh':(left+right)/12,
            'conversion_method':CONFIG['forecast_conversion'],
            'endpoint_rule':np.where(starts<60,'previous_issue_point','current_issue_points')})
        out['endpoint_issue_time'] = pd.NaT
        if anchor is not None:
            out.loc[starts<60,'endpoint_issue_time'] = anchor[0]
        out['date'] = out.interval_start.dt.strftime('%Y-%m-%d')
        out['slot_id'] = out.interval_start.dt.hour*6+out.interval_start.dt.minute//10+1
        frames.append(out)
        # Updating after conversion makes every eligible anchor strictly older.
        for row in group.itertuples():
            available_points[row.target_time] = (issue,row.pv_forecast_kw)
    return pd.concat(frames,ignore_index=True)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--stage',choices=['q1','hourly','all'],default='all')
    args = parser.parse_args()
    if CONFIG['raw_time_label'] != 'interval_end_assumption' or CONFIG['power_interpretation'] != 'interval_average_assumption' or CONFIG['interval_minutes'] != 10:
        raise ValueError('Changed assumptions require corresponding implementation and verification changes')
    q, fixed = prepare_q1()
    if args.stage == 'q1':
        return
    prepare_actual(fixed)
    hourly = prepare_hourly()
    if args.stage == 'hourly':
        return
    derived = convert_forecasts(hourly)
    write(derived,'pv_forecast_10min')
    coverage = derived.groupby('issue_time').agg(interval_count=('interval_start','size'),
        first_interval_start=('interval_start','min'),last_interval_end=('interval_end','max')).reset_index()
    coverage['missing_first_hour'] = coverage.interval_count.eq(138)
    coverage.to_csv(WORK/'data/interim/forecast_coverage.csv',index=False,date_format='%Y-%m-%dT%H:%M:%S')


if __name__ == '__main__':
    main()

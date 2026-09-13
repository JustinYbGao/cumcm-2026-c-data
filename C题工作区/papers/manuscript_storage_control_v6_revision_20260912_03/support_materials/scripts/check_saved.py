"""Lightweight independent checks of saved plans, costs and workbook identities.

This file does not import a production optimizer. The optional full replay uses
prepared original input data but never searches for a new procurement strategy.
"""
from pathlib import Path
from collections import Counter
import argparse
import hashlib
import json
import math
import numpy as np
import pandas as pd
import openpyxl

WORK=Path(__file__).resolve().parents[1]
counts=Counter();errors=[];maximum={}
def read(path):return pd.read_csv(path,float_precision='round_trip')
def check(name,actual,expected,tol=1e-6):
    a,b=np.asarray(actual),np.asarray(expected)
    counts[name]+=int(a.size)
    if a.shape!=b.shape:
        errors.append({'check':name,'shape_actual':list(a.shape),'shape_expected':list(b.shape)});return
    if a.dtype.kind in 'iufc' and b.dtype.kind in 'iufc':
        diff=float(np.max(np.abs(a-b))) if a.size else 0.
        maximum[name]=max(maximum.get(name,0.),diff)
        ok=bool(np.all(np.isfinite(a)) and np.all(np.isfinite(b)) and diff<=tol)
    else:ok=bool(np.array_equal(a,b))
    if not ok:errors.append({'check':name,'max_error':maximum.get(name)})

def replay(frame,initial):
    """Independent sequential balance, reserve discharge and Rule-A cash replay."""
    energy=float(initial);records=[]
    for row in frame.itertuples():
        q=float(row.grid_effective_kwh);q0=float(row.grid_original_kwh)
        reserve=float(row.reserve_threshold_kwh)
        net=q+float(row.pv_actual_kwh)-float(row.load_actual_kwh)
        start=energy
        if net>=0:
            charge=min(net,5000./6,max(0.,(10800.-energy)/.9));discharge=emergency=0.
            surplus=net-charge;energy+=.9*charge
        else:
            charge=surplus=0.
            discharge=min(-net,5000./6,max(0.,.9*(energy-reserve)))
            emergency=-net-discharge;energy-=discharge/.9
        price=float(row.price_yuan_per_kwh)
        original=price*q0;increase=1.5*price*max(q-q0,0.)
        decrease=-.5*price*max(q0-q,0.)
        contract=original+increase+decrease;urgent=5*price*emergency
        records.append({'charge_actual_kwh':charge,'discharge_actual_kwh':discharge,
            'emergency_kwh':emergency,'surplus_kwh':surplus,
            'unused_grid_kwh':min(q,surplus),'pv_curtailment_kwh':surplus-min(q,surplus),
            'energy_start_actual_kwh':start,'energy_end_actual_kwh':energy,
            'original_cost_yuan':original,'increase_cost_yuan':increase,
            'decrease_adjustment_yuan':decrease,'contract_cost_yuan':contract,
            'emergency_cost_yuan':urgent,'total_cost_yuan':contract+urgent})
    return pd.DataFrame(records,index=frame.index)

def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--full-inputs',action='store_true',help='Replay all selected annual paths using data prepared from the original attachments')
    parser.add_argument('--output',type=Path,default=WORK/'check_results.json')
    args=parser.parse_args()
    mapping=json.loads((WORK/'workbooks/policy_mapping.json').read_text())
    workbook_records=[]
    for item in mapping['workbooks']:
        path=WORK/'workbooks'/item['file']
        check('workbook_binary_identity',hashlib.sha256(path.read_bytes()).hexdigest(),item['english_workbook_sha256'])
        wb=openpyxl.load_workbook(path,data_only=False)
        check('sheet_names',wb.sheetnames,[s['english_sheet_name'] for s in item['sheets']])
        cell_count=formulas=dates=0
        for sheet,expected in zip(wb,item['sheets']):
            check('worksheet_dimensions',[sheet.max_row,sheet.max_column],[expected['rows'],expected['columns']])
            for row in sheet:
                for cell in row:
                    cell_count+=1;formulas+=cell.data_type=='f';dates+=cell.data_type=='d'
                    if cell.data_type=='e':errors.append({'check':'spreadsheet_error','file':item['file'],'cell':cell.coordinate})
        workbook_records.append({'file':item['file'],'cells_read':cell_count,'formulas':formulas,'dates':dates})
        wb.close()
    q1=read(WORK/'saved/q1/baseline/schedule.csv')
    check('q1_slot_keys',q1.slot_id,np.arange(1,145))
    check('q1_balance',q1.grid_kwh+q1.pv_forecast_kwh+q1.discharge_kwh-q1.load_kwh-q1.charge_kwh-q1.curtailment_kwh,np.zeros(144))
    check('q1_state',q1.energy_start_kwh+.9*q1.charge_kwh-q1.discharge_kwh/.9,q1.energy_end_kwh)
    check('q1_state_continuity',q1.energy_start_kwh.iloc[1:],q1.energy_end_kwh.iloc[:-1])
    check('q1_endpoints',[q1.energy_start_kwh.iloc[0],q1.energy_end_kwh.iloc[-1]],[6000.,6000.])
    check('q1_cash',q1.grid_kwh*q1.price_yuan_per_kwh,q1.cost_yuan)
    check('q1_modes',np.minimum(q1.charge_kwh,q1.discharge_kwh),np.zeros(144))
    summary=json.loads((WORK/'saved/q1/baseline/summary.json').read_text())
    check('q1_total',math.fsum(q1.cost_yuan),summary['cost_yuan'])
    fixed=read(WORK/'saved/fixed_price.csv').price_yuan_per_kwh.to_numpy()
    check('diagnostic_inventory_value',float(fixed.mean()*.9),.6895775,1e-12)
    warm=read(WORK/'saved/warmup/ledger.csv')
    check('warmup_start',float(warm.energy_start_actual_kwh.iloc[0]),6000.)
    check('warmup_continuity',warm.energy_start_actual_kwh.iloc[1:],warm.energy_end_actual_kwh.iloc[:-1])
    check('warmup_final',float(warm.energy_end_actual_kwh.iloc[-1]),7268.4231640740745)

    config=json.loads((WORK/'configs/storage_control_v6/experiments.json').read_text())['policies']
    selection=json.loads((WORK/'saved/comparisons/selection.json').read_text())['selected']
    full_source=read(WORK/'data/processed/actual_10min.csv') if args.full_inputs else None
    for name in [value for key,value in selection.items() if key!='q1']:
        folder=WORK/'saved/selected'/name
        with np.load(folder/'execution_plans.npz',allow_pickle=False) as packed:
            plans=pd.DataFrame({key:packed[key] for key in packed.files})
        with np.load(folder/'plan_versions.npz',allow_pickle=False) as packed:
            versions=pd.DataFrame({key:packed[key] for key in packed.files})
        check('annual_row_count',len(plans),48096,0)
        check('annual_slot_keys',plans.slot_id,np.tile(np.arange(1,145),334),0)
        check('annual_dates',plans.date,np.repeat(pd.date_range('2025-02-01','2025-12-31').strftime('%Y-%m-%d'),144))
        check('q_nonnegative',bool((plans[['grid_original_kwh','grid_effective_kwh']]>=0).all().all()),True)
        check('reserve_bounds',bool(plans.reserve_threshold_kwh.between(1200.,10800.).all()),True)
        hours=[0,6,12,18] if config[name]['updates'] else [0]
        for date,day in plans.groupby('date',sort=False):
            current_q=current_r=None
            for hour in hours:
                issued=versions.loc[(versions.date==date)&(versions.issue_hour==hour)]
                start=hour*6
                check('version_slot_keys',issued.slot_id,np.arange(start+1,145),0)
                check('version_q0',issued.grid_original_kwh,day.grid_original_kwh.to_numpy()[start:],0)
                if hour==0:
                    current_q=issued.grid_kwh.to_numpy().copy();current_r=issued.reserve_threshold_kwh.to_numpy().copy()
                else:
                    current_q[start:]=issued.grid_kwh.to_numpy();current_r[start:]=issued.reserve_threshold_kwh.to_numpy()
                end=min(start+36,144) if config[name]['updates'] else 144
                check('executed_q_from_issued_version',day.grid_effective_kwh.to_numpy()[start:end],current_q[start:end],0)
                check('executed_R_from_issued_version',day.reserve_threshold_kwh.to_numpy()[start:end],current_r[start:end],0)
        four=read(folder/'ledger_four_dates.csv')
        for date,frame in four.groupby('date',sort=False):
            rebuilt=replay(frame,float(frame.energy_start_actual_kwh.iloc[0]))
            for column in rebuilt:check('four_date_'+column,rebuilt[column],frame[column])
            check('four_date_energy_bounds',bool(rebuilt.energy_end_actual_kwh.between(1200.-1e-6,10800.+1e-6).all()),True)
            check('four_date_modes',np.minimum(rebuilt.charge_actual_kwh,rebuilt.discharge_actual_kwh),np.zeros(144))
        if args.full_inputs:
            source=full_source.loc[full_source.date>='2025-02-01'].reset_index(drop=True)
            frame=plans.copy()
            check('raw_actual_keys',source[['date','slot_id']].astype(str),frame[['date','slot_id']].astype(str))
            frame['load_actual_kwh']=source.load_actual_kwh;frame['pv_actual_kwh']=source.pv_actual_kwh
            frame['price_yuan_per_kwh']=np.tile(fixed,334) if config[name]['billing']=='fixed' else source.actual_price_yuan_per_kwh
            rebuilt=replay(frame,7268.4231640740745)
            saved=read(WORK/f'saved/all_policies/{name}/daily.csv')
            rebuilt['date']=frame.date
            sums=rebuilt.groupby('date').sum(numeric_only=True)
            for column in ['original_cost_yuan','increase_cost_yuan','decrease_adjustment_yuan','contract_cost_yuan','emergency_cost_yuan','total_cost_yuan','charge_actual_kwh','discharge_actual_kwh','emergency_kwh']:
                check('annual_'+column,sums[column],saved[column],1e-5)
            check('annual_end_energy',rebuilt.groupby('date').energy_end_actual_kwh.last(),saved.energy_end_actual_kwh)
    for name in config:
        folder=WORK/'saved/all_policies'/name
        daily=read(folder/'daily.csv');summary=json.loads((folder/'summary.json').read_text())
        check('all_policy_days',len(daily),334,0)
        for field in ['original_cost_yuan','increase_cost_yuan','decrease_adjustment_yuan','contract_cost_yuan','emergency_cost_yuan','total_cost_yuan']:
            check('all_policy_'+field,math.fsum(daily[field]),summary[field],1e-5)
        check('all_policy_cash_components',daily.contract_cost_yuan+daily.emergency_cost_yuan,daily.total_cost_yuan,1e-6)
    report={'passed':not errors,'full_input_replay':args.full_inputs,'annual_optimizations_run':0,
            'checks':dict(counts),'maximum_absolute_errors':maximum,'errors':errors,'workbooks':workbook_records,
            'limitations':'Default mode checks four specified dates, all annual issued q/R versions, all saved daily summaries and exact workbook identities. It does not certify optimality or rerun historical policy searches.'}
    args.output.parent.mkdir(parents=True,exist_ok=True)
    args.output.write_text(json.dumps(report,indent=2)+'\n')
    print(json.dumps({'passed':report['passed'],'checks':sum(counts.values()),'errors':len(errors),'full_input_replay':args.full_inputs}))
    raise SystemExit(0 if report['passed'] else 1)

if __name__=='__main__':main()

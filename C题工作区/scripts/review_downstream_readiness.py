"""Read-only downstream readiness audit; writes only its own evidence report."""
from pathlib import Path
import json
import numpy as np
import pandas as pd
import openpyxl
from prepare_data import WORK, OUT, SOURCE, history_at, forecasts_at
from diagnose_forecasts import errors_at

def main():
    q=pd.read_csv(OUT/'q1_day.csv')
    a=pd.read_csv(OUT/'actual_10min.csv',parse_dates=['interval_start','interval_end','available_time'])
    f=pd.read_csv(OUT/'pv_forecast_10min.csv',parse_dates=['issue_time','interval_start','interval_end'])
    e=pd.read_csv(OUT/'pv_forecast_hourly_evaluation.csv',parse_dates=['issue_time','target_time','error_available_time'],dtype={'actual_positive_pv':'boolean'})
    config=json.loads((WORK/'configs/data_processing.json').read_text())
    checks={}
    counts=f.loc[f.issue_time.ge('2025-02-01')].groupby('issue_time').size()
    expected=pd.date_range('2025-02-01',periods=334*4,freq='6h')
    checks['feb_dec_all_1336_releases_have_144_intervals']=bool(counts.index.equals(expected) and counts.eq(144).all())
    checks['q1_finite_required_numeric']=bool(np.isfinite(q[['price_yuan_per_kwh','load_kwh','pv_forecast_kwh']].to_numpy()).all())
    checks['q1_nonnegative_price_supply_demand']=bool(q[['price_yuan_per_kwh','load_kwh','pv_forecast_kwh']].ge(0).all().all())
    # Exercise the documented downstream selection, using CSV readback.
    picks={}
    for decision in ['2025-02-01 00:00','2025-03-20 06:00','2025-12-31 18:00']:
        now=pd.Timestamp(decision)
        usable=forecasts_at(f,now)
        exact=usable.loc[usable.issue_time.eq(now)]
        checks['current_release_full_24h_'+decision]=len(exact)==144 and exact.interval_start.min()==now and exact.interval_end.max()==now+pd.Timedelta(hours=24)
        old=history_at(a,now);err=errors_at(e,now)
        checks['causal_'+decision]=bool(old.available_time.le(now).all() and err.error_available_time.le(now).all())
        picks[decision]={'actual_rows':len(old),'error_rows':len(err),'current_release_rows':len(exact)}
    valid=e.loc[e.comparison_status.eq('comparable')]
    checks['error_sign_is_forecast_minus_actual']=bool(np.allclose(valid.error_kwh,valid.forecast_hour_kwh-valid.actual_hour_kwh))
    wb=openpyxl.load_workbook(SOURCE/'附件5/result1.xlsx',read_only=True,data_only=True)
    ws=wb['计划购电量']
    template={'first_row_label':ws['A2'].value,'last_row_label':ws['A145'].value}
    wb.close()
    report={'checks':checks,'passed':all(checks.values()),'assumptions':config,'source_template_q1':template,
            'q1_intervals':{'first':[q.start_time.iloc[0],q.end_time.iloc[0]],'last':[q.start_time.iloc[-1],q.end_time.iloc[-1]]},
            'decision_examples':picks,
            'notes':['Under current documented assumptions only; not official resolution of time-label ambiguity.',
                     'Existing evaluation error is forecast-actual. A deficit-risk residual actual-forecast requires sign reversal.',
                     'No prediction or scheduling model has been trained/solved by this audit.']}
    (WORK/'reports/review_downstream_readiness.json').write_text(json.dumps(report,ensure_ascii=False,indent=2))
    print(json.dumps(report,ensure_ascii=False,indent=2))
    if not report['passed']:raise SystemExit(1)
if __name__=='__main__':main()

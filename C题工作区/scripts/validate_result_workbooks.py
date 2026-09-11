#!/usr/bin/env python3
"""Independent, read-only validation of the five review result workbooks."""
from __future__ import annotations
import argparse,csv,datetime as dt,hashlib,json,math,sys,unittest
from collections import defaultdict
from pathlib import Path
import openpyxl
ROOT=Path(__file__).resolve().parents[1]; REPO=ROOT.parent; OUT=ROOT/'outputs/revision_v1'; REPORT=ROOT/'reports/revision_v1'; TEMPLATES=REPO/'CUMCM2026Problems/C题/附件/附件5'
TOL=1e-6; BAD={'#NULL!','#DIV/0!','#VALUE!','#REF!','#NAME?','#NUM!','#N/A','#GETTING_DATA','…','⁝','...'}
CONFIG={
'result1.xlsx':('results/q1/baseline/schedule.csv',['计划购电量','充放电量'],'q1'),
'result2.xlsx':('results/q2/selected/ledger.csv',['计划购电量','充放电量','紧急购电量'],'basic'),
'result3.xlsx':('results/robustness/fixed_w28/ledger.csv',['计划购电量','调整购电量','充放电量','紧急购电量'],'adjusted'),
'result4-2.xlsx':('results/q4/q42_ols/ledger.csv',['计划购电量','充放电量','紧急购电量'],'basic'),
'result4-3.xlsx':('results/robustness/variable_w28/ledger.csv',['计划购电量','调整购电量','充放电量','紧急购电量'],'adjusted')}
def sha256(p):
 h=hashlib.sha256()
 with p.open('rb') as f:
  for b in iter(lambda:f.read(1<<20),b''): h.update(b)
 return h.hexdigest()
def fmt(m): return '24:00' if m==1440 else f'{m//60:02d}:{m%60:02d}'
def interval_label(s): return f'{fmt((s-1)*10)}-{fmt(s*10)}'
def interval_labels(): return [interval_label(s) for s in range(1,145)]
def read_csv(p):
 with p.open(encoding='utf-8-sig',newline='') as f:return list(csv.DictReader(f))
def num(r,k): return float(r[k])
def dtext(v): return v.strftime('%Y-%m-%d') if isinstance(v,(dt.date,dt.datetime)) else str(v)[:10]
def emergency_runs(rows):
 by=defaultdict(list)
 for r in rows:
  if num(r,'emergency_kwh')>1e-7: by[r['date']].append(r)
 out=[]
 for date in sorted(by):
  active=sorted(by[date],key=lambda r:int(r['slot_id'])); group=[active[0]]
  for r in active[1:]:
   if int(r['slot_id'])==int(group[-1]['slot_id'])+1: group.append(r)
   else: out.append(finish_run(date,group)); group=[r]
  out.append(finish_run(date,group))
 return out
def finish_run(date,g): return date,f'{fmt((int(g[0]["slot_id"])-1)*10)}-{fmt(int(g[-1]["slot_id"])*10)}',sum(num(r,'emergency_kwh') for r in g)
def color_value(c):
 if c is None:return None
 return (c.type,c.rgb if c.type=='rgb' else None,c.indexed if c.type=='indexed' else None,c.theme if c.type=='theme' else None,float(c.tint or 0),bool(c.auto))
def semantic_style(cell):
 f=cell.font;fill=cell.fill;b=cell.border;al=cell.alignment;side=lambda s:(s.style,color_value(s.color)) if s is not None else (None,None)
 return ((f.name,float(f.sz) if f.sz is not None else None,bool(f.bold),bool(f.italic),f.underline,color_value(f.color)),
         (fill.fill_type,color_value(fill.fgColor),color_value(fill.bgColor)),
         tuple(side(getattr(b,k)) for k in ('left','right','top','bottom','diagonal','vertical','horizontal')),
         (al.horizontal,al.vertical,bool(al.wrap_text)),cell.number_format)
def parse_stamp(v): return dt.datetime.fromisoformat(str(v).replace('T',' '))
def validate_ledger_keys(a,rows,name,actual):
 seen=set()
 for r in rows:
  date=r['date'];slot=int(r['slot_id']);key=(date,slot);base=dt.datetime.fromisoformat(date);start=base+dt.timedelta(minutes=(slot-1)*10);end=base+dt.timedelta(minutes=slot*10)
  a.check(key not in seen,'duplicate_ledger_key',f'{name}/ledger {date}#{slot}');seen.add(key)
  a.check(parse_stamp(r['interval_start'])==start,'interval_start_alignment',f'{name}/ledger {date}#{slot}',start.isoformat(' '),r['interval_start'])
  a.check(parse_stamp(r['interval_end'])==end,'interval_end_alignment',f'{name}/ledger {date}#{slot}',end.isoformat(' '),r['interval_end'])
  if key in actual:
   src=actual[key]
   if 'load_actual_kwh' in r:a.close(num(r,'load_actual_kwh'),num(src,'load_actual_kwh'),f'{name}/ledger load {date}#{slot}')
   if 'pv_actual_kwh' in r:a.close(num(r,'pv_actual_kwh'),num(src,'pv_actual_kwh'),f'{name}/ledger PV {date}#{slot}')
 a.check(len(seen)==len(rows),'ledger_key_coverage',f'{name}/ledger')
def independent_prices(rows,kind,fixed,actual):
 prices={}; emergency={}
 for r in rows:
  key=(r['date'],int(r['slot_id'])); is_fixed=kind in ('q1','basic-fixed','adjusted-fixed')
  prices[key]=num(fixed[int(r['slot_id'])],'price_yuan_per_kwh') if is_fixed else num(actual[key],'actual_price_yuan_per_kwh'); emergency[key]=prices[key]
 return prices,emergency
class Audit:
 def __init__(self):self.errors=[];self.checks=0;self.numeric=0
 def check(self,ok,code,where,expected=None,actual=None):
  self.checks+=1
  if not ok:
   e={'code':code,'where':where}
   if expected is not None:e['expected']=expected
   if actual is not None:e['actual']=actual
   self.errors.append(e)
 def close(self,actual,expected,where):
  self.numeric+=1
  try:ok=math.isfinite(float(actual)) and math.isclose(float(actual),float(expected),rel_tol=1e-9,abs_tol=TOL)
  except (TypeError,ValueError):ok=False
  self.check(ok,'numeric_mismatch',where,expected,actual)
def validate_plan(a,wf,wd,rows,adjusted,name,prices):
 dates=sorted({r['date'] for r in rows}); by={(r['date'],int(r['slot_id'])):r for r in rows}
 a.check((wf.max_row,wf.max_column)==(335,147),'plan_shape',f'{name}/{wf.title}',[335,147],[wf.max_row,wf.max_column])
 a.check([wf.cell(1,c).value for c in range(2,146)]==interval_labels(),'interval_headers',f'{name}/{wf.title}')
 a.check([dtext(wf.cell(r,1).value) for r in range(2,336)]==dates,'date_coverage',f'{name}/{wf.title}')
 for ri,date in enumerate(dates,2):
  day=[by[(date,s)] for s in range(1,145)]; key='grid_effective_kwh' if adjusted else ('grid_plan_kwh' if 'grid_plan_kwh' in day[0] else 'grid_original_kwh')
  for s,r in enumerate(day,2):a.close(wf.cell(ri,s).value,num(r,key),f'{name}/{wf.title}!{wf.cell(ri,s).coordinate}')
  a.check(wf.cell(ri,146).value==f'=SUM(B{ri}:EO{ri})','sum_formula',f'{name}/{wf.title}!EP{ri}')
  a.close(wd.cell(ri,146).value,sum(num(r,key) for r in day),f'{name}/{wf.title}!EP{ri}[cache]')
  if adjusted:
   fee=sum(prices[(date,int(r['slot_id']))]*(num(r,'grid_original_kwh')+1.5*max(num(r,key)-num(r,'grid_original_kwh'),0)-.5*max(num(r,'grid_original_kwh')-num(r,key),0)) for r in day)
  else: fee=sum(prices[(date,int(r['slot_id']))]*float(wf.cell(ri,s).value) for s,r in enumerate(day,2))
  a.close(wf.cell(ri,147).value,fee,f'{name}/{wf.title}!EQ{ri}')
def validate_battery(a,ws,rows,name):
 dates=sorted({r['date'] for r in rows}); by={(r['date'],int(r['slot_id'])):r for r in rows}; periods=[f'{h:02d}:00-{h+4:02d}:00' for h in range(0,24,4)]
 a.check((ws.max_row,ws.max_column)==(2005,6),'battery_shape',f'{name}/充放电量',[2005,6],[ws.max_row,ws.max_column])
 for di,date in enumerate(dates):
  for b in range(6):
   rr=2+di*6+b; chunk=[by[(date,s)] for s in range(b*24+1,b*24+25)]
   a.check(dtext(ws.cell(rr,1).value)==date,'battery_date',f'{name}/充放电量!A{rr}')
   a.check(ws.cell(rr,2).value==periods[b],'battery_period',f'{name}/充放电量!B{rr}',periods[b],ws.cell(rr,2).value)
   a.close(ws.cell(rr,3).value,sum(num(r,'charge_actual_kwh') for r in chunk),f'{name}/充放电量!C{rr}'); a.close(ws.cell(rr,4).value,sum(num(r,'discharge_actual_kwh') for r in chunk),f'{name}/充放电量!D{rr}')
   mark='00:00' if b==0 else '24:00' if b==1 else None; soc=num(by[(date,1)],'energy_start_actual_kwh') if b==0 else num(by[(date,144)],'energy_end_actual_kwh') if b==1 else None
   a.check(ws.cell(rr,5).value==mark,'soc_mark',f'{name}/充放电量!E{rr}',mark,ws.cell(rr,5).value)
   if soc is None:a.check(ws.cell(rr,6).value is None,'soc_extra',f'{name}/充放电量!F{rr}')
   else:a.close(ws.cell(rr,6).value,soc,f'{name}/充放电量!F{rr}')
 for r in rows:
  end=num(r,'energy_start_actual_kwh')+.9*num(r,'charge_actual_kwh')-num(r,'discharge_actual_kwh')/.9
  a.close(num(r,'energy_end_actual_kwh'),end,f'{name}/ledger SOC {r["date"]}#{r["slot_id"]}')
  a.check(not(num(r,'charge_actual_kwh')>TOL and num(r,'discharge_actual_kwh')>TOL),'simultaneous_charge_discharge',f'{name}/ledger {r["date"]}#{r["slot_id"]}')
def validate_events(a,ws,rows,name):
 exp=emergency_runs(rows); act=[(dtext(ws.cell(r,1).value),ws.cell(r,2).value,ws.cell(r,3).value) for r in range(2,ws.max_row+1)]
 a.check(len(act)==len(exp),'event_count',f'{name}/紧急购电量',len(exp),len(act))
 for i,e in enumerate(exp[:len(act)]):a.check(act[i][:2]==e[:2],'event_key',f'{name}/紧急购电量 row {i+2}',e[:2],act[i][:2]);a.close(act[i][2],e[2],f'{name}/紧急购电量!C{i+2}')
def validate_q1(a,wb,rows,name):
 ws=wb['计划购电量']; by={int(r['slot_id']):r for r in rows}; periods=[f'{h:02d}:00-{h+4:02d}:00' for h in range(0,24,4)]
 a.check((ws.max_row,ws.max_column)==(145,2),'q1_plan_shape',f'{name}/计划购电量');a.check([ws.cell(r,1).value for r in range(2,146)]==interval_labels(),'interval_headers',f'{name}/计划购电量')
 for s in range(1,145):a.close(ws.cell(s+1,2).value,num(by[s],'grid_kwh'),f'{name}/计划购电量!B{s+1}')
 bs=wb['充放电量']
 for b in range(6):
  chunk=[by[s] for s in range(b*24+1,b*24+25)];a.check(bs.cell(b+2,1).value==periods[b],'q1_battery_period',f'{name}/充放电量!A{b+2}')
  a.close(bs.cell(b+2,2).value,sum(num(r,'charge_kwh') for r in chunk),f'{name}/充放电量!B{b+2}');a.close(bs.cell(b+2,3).value,sum(num(r,'discharge_kwh') for r in chunk),f'{name}/充放电量!C{b+2}')
 a.close(bs.cell(2,5).value,num(by[1],'energy_start_kwh'),f'{name}/充放电量!E2');a.close(bs.cell(3,5).value,num(by[144],'energy_end_kwh'),f'{name}/充放电量!E3')
def validate_content(a,wb,template,name):
 for ws in wb:
  for row in ws.iter_rows():
   for cell in row:
    if isinstance(cell.value,str):a.check(cell.value.strip() not in BAD,'forbidden_cell',f'{name}/{ws.title}!{cell.coordinate}',actual=cell.value)
  a.check(semantic_style(ws['A1'])==semantic_style(template[ws.title]['A1']),'header_style',f'{name}/{ws.title}!A1')
  a.check(ws.max_row>=template[ws.title].max_row and ws.max_column>=template[ws.title].max_column,'template_not_preserved',f'{name}/{ws.title}')
def validate_mapping(a):
 rows=read_csv(OUT/'template_time_mapping.csv'); expected_sheets={n:[s for s in sheets if s in ('计划购电量','调整购电量')] for n,(_,sheets,_) in CONFIG.items()}; seen=set()
 a.check(len(rows)==1008,'mapping_count','template_time_mapping.csv',1008,len(rows))
 for r in rows:
  name=r['file'];sheet=r.get('sheet','计划购电量');slot=int(r['source_slot_id']);key=(name,sheet,slot);seen.add(key);template=openpyxl.load_workbook(TEMPLATES/name,read_only=True,data_only=False)
  expected_cell=f'A{slot+1}' if name=='result1.xlsx' else openpyxl.utils.get_column_letter(slot+1)+'1'
  a.check(sheet in expected_sheets.get(name,[]),'mapping_sheet',f'{name}/{sheet}')
  a.check(r['template_cell']==expected_cell,'mapping_cell',f'{name}/{sheet}#{slot}',expected_cell,r['template_cell'])
  a.check(str(template[sheet][expected_cell].value)==r['original_label'],'mapping_original_label',f'{name}/{sheet}!{expected_cell}',template[sheet][expected_cell].value,r['original_label'])
  a.check(r['review_label']==interval_label(slot),'mapping_review_label',f'{name}/{sheet}#{slot}',interval_label(slot),r['review_label'])
  a.check((int(r['start_minute']),int(r['end_minute']))==((slot-1)*10,slot*10),'mapping_minutes',f'{name}/{sheet}#{slot}')
  a.check(r['operation']=='relabel_copy_only_no_data_shift','mapping_operation',f'{name}/{sheet}#{slot}')
 expected={(n,s,slot) for n,sheets in expected_sheets.items() for s in sheets for slot in range(1,145)}
 a.check(seen==expected,'mapping_coverage','template_time_mapping.csv',len(expected),len(seen))
def validate_all():
 a=Audit(); recorded=json.loads((OUT/'source_hashes.json').read_text()) if (OUT/'source_hashes.json').exists() else {}; hashes={}; summary={r['file']:r for r in read_csv(OUT/'cost_summary.csv')}
 fixed_path=ROOT/'data/processed/fixed_price.csv';actual_path=ROOT/'data/processed/actual_10min.csv';fixed={int(r['slot_id']):r for r in read_csv(fixed_path)}; actual={(r['date'],int(r['slot_id'])):r for r in read_csv(actual_path)}
 for p in (fixed_path,actual_path):hashes[str(p.relative_to(REPO))]=sha256(p)
 for name,(rel,sheets,kind) in CONFIG.items():
  lp,tp,op=ROOT/rel,TEMPLATES/name,OUT/name; rows=read_csv(lp)
  for p in (lp,tp):
   key=str(p.relative_to(REPO));hashes[key]=sha256(p);a.check(recorded.get(key)==hashes[key],'source_hash_changed',key,recorded.get(key),hashes[key])
  wb=openpyxl.load_workbook(op,data_only=False);wd=openpyxl.load_workbook(op,data_only=True);template=openpyxl.load_workbook(tp)
  if kind!='q1':validate_ledger_keys(a,rows,name,actual)
  a.check(wb.sheetnames==sheets,'sheet_order',name,sheets,wb.sheetnames);validate_content(a,wb,template,name)
  price_kind='q1' if kind=='q1' else 'adjusted-fixed' if name=='result3.xlsx' else 'basic-fixed' if name=='result2.xlsx' else 'variable'
  prices,emergency_prices=({}, {}) if kind=='q1' else independent_prices(rows,price_kind,fixed,actual)
  if kind=='q1':validate_q1(a,wb,rows,name)
  else:
   validate_plan(a,wb['计划购电量'],wd['计划购电量'],rows,False,name,prices)
   if kind=='adjusted':validate_plan(a,wb['调整购电量'],wd['调整购电量'],rows,True,name,prices)
   validate_battery(a,wb['充放电量'],rows,name);validate_events(a,wb['紧急购电量'],rows,name)
  if kind=='q1':contract=sum(float(wb['计划购电量'].cell(s+1,2).value)*num(fixed[s],'price_yuan_per_kwh') for s in range(1,145));emergency=0
  else:
   fee_sheet=wb['调整购电量'] if kind=='adjusted' else wb['计划购电量'];contract=sum(float(fee_sheet.cell(r,147).value) for r in range(2,336));emergency=sum(5*emergency_prices[(r['date'],int(r['slot_id']))]*num(r,'emergency_kwh') for r in rows)
  sr=summary.get(name);a.check(sr is not None,'summary_missing',name)
  if sr:a.close(sr['contract_cost_yuan'],contract,f'cost_summary/{name}/contract');a.close(sr['emergency_cost_yuan'],emergency,f'cost_summary/{name}/emergency');a.close(sr['total_cost_yuan'],contract+emergency,f'cost_summary/{name}/total');a.check(int(sr['days'])==(1 if kind=='q1' else 334),'summary_days',f'cost_summary/{name}')
 validate_mapping(a)
 return {'status':'pass' if not a.errors else 'fail','checks':a.checks,'numeric_cells_checked':a.numeric,'error_count':len(a.errors),'errors':a.errors[:200],'source_hashes':hashes,'validator_sha256':sha256(Path(__file__)),'independence':'Expected values derived from original templates and CSV ledgers; exporter/payload not imported or read.'}
def write_reports(r):
 REPORT.mkdir(parents=True,exist_ok=True);(REPORT/'export_validation.json').write_text(json.dumps(r,ensure_ascii=False,indent=2)+'\n')
 lines=['# Result workbook independent validation','',f'- Status: **{r["status"].upper()}**',f'- Checks: {r["checks"]}',f'- Numeric comparisons: {r["numeric_cells_checked"]}',f'- Errors: {r["error_count"]}',f'- Validator SHA-256: `{r["validator_sha256"]}`','',r['independence']]
 if r['errors']:lines+=['','## Errors','']+[f'- `{e["code"]}` at `{e["where"]}`: expected `{e.get("expected")}`, actual `{e.get("actual")}`' for e in r['errors']]
 (REPORT/'export_validation.md').write_text('\n'.join(lines)+'\n')
class UnitTests(unittest.TestCase):
 def test_labels(self):self.assertEqual((interval_labels()[0],interval_labels()[-1],len(interval_labels())),('00:00-00:10','23:50-24:00',144))
 def test_runs_dates(self):
  rows=[{'date':'2025-02-01','slot_id':'144','emergency_kwh':'2'},{'date':'2025-02-02','slot_id':'1','emergency_kwh':'3'}];self.assertEqual(emergency_runs(rows),[('2025-02-01','23:50-24:00',2),('2025-02-02','00:00-00:10',3)])
 def test_semantic_style_ignores_registry_metadata(self):
  from openpyxl.styles import Font
  w=openpyxl.Workbook();a=w.active['A1'];b=w.active['B1'];c=w.active['C1'];a.font=Font(name='宋体',size=10);b.font=Font(name='宋体',size=10,charset=134);c.font=Font(name='宋体',size=10,bold=True)
  self.assertNotEqual(a.style_id,b.style_id);self.assertEqual(semantic_style(a),semantic_style(b));self.assertNotEqual(semantic_style(a),semantic_style(c))
def main():
 p=argparse.ArgumentParser();p.add_argument('--self-test',action='store_true');args=p.parse_args()
 if args.self_test:return 0 if unittest.TextTestRunner().run(unittest.defaultTestLoader.loadTestsFromTestCase(UnitTests)).wasSuccessful() else 1
 r=validate_all();write_reports(r);print(json.dumps({k:r[k] for k in ('status','checks','numeric_cells_checked','error_count')}));return 0 if r['status']=='pass' else 1
if __name__=='__main__':sys.exit(main())

#!/usr/bin/env python3
"""Prepare the five selected workbook exports and paper-table extracts.

This module writes JSON/CSV/Markdown inputs for the Artifact Tool builder.  It
never writes an XLSX file.  Selection is read only from the frozen v5 result.
"""
from __future__ import annotations

import csv
import hashlib
import json
from contextlib import closing
from datetime import datetime, timedelta
from pathlib import Path

import openpyxl


ROOT = Path(__file__).resolve().parents[2]
REPO = ROOT.parent
OUT = ROOT / "outputs/unified_direct_v5"
REPORT = ROOT / "reports/unified_direct_v5"
SELECTION = ROOT / "results/unified_direct_v5/selection.json"
TEMPLATES = REPO / "CUMCM2026Problems/C题/附件/附件5"
SPECIFIED_DATES = ("2025-03-20", "2025-06-21", "2025-09-23", "2025-12-21")


def read_csv(path: Path) -> list[dict[str, str]]:
    with path.open(encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1 << 20), b""):
            digest.update(block)
    return digest.hexdigest()


def rel(path: Path) -> str:
    return str(path.relative_to(ROOT)) if path.is_relative_to(ROOT) else str(path.relative_to(REPO))


def number(row: dict[str, str], key: str) -> float:
    return float(row[key])


def clock(minute: int) -> str:
    return "24:00" if minute == 1440 else f"{minute // 60:02d}:{minute % 60:02d}"


def period(start: int, end: int) -> str:
    return f"{clock(start)}-{clock(end)}"


def interval_labels() -> list[str]:
    return [period(i * 10, (i + 1) * 10) for i in range(144)]


def write_csv(path: Path, records: list[dict]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(records[0]))
        writer.writeheader()
        writer.writerows(records)


def md_value(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float):
        return f"{value:.6f}"
    return str(value).replace("|", "\\|")


def write_markdown(path: Path, title: str, records: list[dict]) -> None:
    fields = list(records[0])
    lines = [f"# {title}", "", "| " + " | ".join(fields) + " |", "| " + " | ".join("---" for _ in fields) + " |"]
    lines += ["| " + " | ".join(md_value(row[field]) for field in fields) + " |" for row in records]
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def keyed_year(rows: list[dict[str, str]]) -> tuple[list[str], dict[tuple[str, int], dict[str, str]]]:
    by_key: dict[tuple[str, int], dict[str, str]] = {}
    for row in rows:
        date = row["date"]
        slot = int(row["slot_id"])
        start = datetime.fromisoformat(row["interval_start"])
        end = datetime.fromisoformat(row["interval_end"])
        expected_start = datetime.fromisoformat(date) + timedelta(minutes=(slot - 1) * 10)
        if start != expected_start or end != start + timedelta(minutes=10):
            raise ValueError(f"misaligned ledger interval: {date} slot {slot}")
        key = (date, slot)
        if key in by_key:
            raise ValueError(f"duplicate ledger interval: {key}")
        by_key[key] = row
    dates = [(datetime(2025, 2, 1) + timedelta(days=i)).date().isoformat() for i in range(334)]
    expected = {(date, slot) for date in dates for slot in range(1, 145)}
    if set(by_key) != expected:
        raise ValueError(f"ledger coverage mismatch: expected {len(expected)}, got {len(by_key)}")
    return dates, by_key


def emergency_events(day_rows: list[dict[str, str]]) -> list[dict]:
    events: list[dict] = []
    i = 0
    while i < 144:
        if number(day_rows[i], "emergency_kwh") <= 1e-7:
            i += 1
            continue
        start = i
        energy = 0.0
        while i < 144 and number(day_rows[i], "emergency_kwh") > 1e-7:
            energy += number(day_rows[i], "emergency_kwh")
            i += 1
        events.append({"interval": period(start * 10, i * 10), "emergency_kwh": energy})
    return events


def component_sums(rows: list[dict[str, str]], q1: bool, q2: bool) -> dict[str, float]:
    if q1:
        total = sum(number(row, "cost_yuan") for row in rows)
        return {
            "original_cost_yuan": total,
            "increase_cost_yuan": 0.0,
            "decrease_adjustment_yuan": 0.0,
            "contract_cost_yuan": total,
            "emergency_cost_yuan": 0.0,
            "total_cost_yuan": total,
        }
    original_key = "planned_cost_yuan" if q2 else "original_cost_yuan"
    contract_key = "planned_cost_yuan" if q2 else "contract_cost_yuan"
    return {
        "original_cost_yuan": sum(number(row, original_key) for row in rows),
        "increase_cost_yuan": 0.0 if q2 else sum(number(row, "increase_cost_yuan") for row in rows),
        "decrease_adjustment_yuan": 0.0 if q2 else sum(number(row, "decrease_adjustment_yuan") for row in rows),
        "contract_cost_yuan": sum(number(row, contract_key) for row in rows),
        "emergency_cost_yuan": sum(number(row, "emergency_cost_yuan") for row in rows),
        "total_cost_yuan": sum(number(row, "total_cost_yuan") for row in rows),
    }


def hashes_for(name: str, ledger: Path, template: Path, strategy: str) -> dict[str, str]:
    if name == "result1":
        paths = [ledger, template, ROOT / "scripts/solve_q1.py", ROOT / "results/q1/baseline/config_snapshot.json"]
    elif name == "result2":
        paths = [ledger, template, ROOT / "scripts/q2_direct_v4/run_experiments.py", ROOT / "scripts/q2_direct_v4/base_policy.py", ROOT / "results/q2_direct_v4/runs/D112/config.json"]
    else:
        paths = [
            ledger,
            template,
            ROOT / f"results/unified_direct_v5/runs/{strategy}/config.json",
            ROOT / "configs/unified_direct_v5/experiments.json",
            ROOT / "reports/unified_direct_v5/protocol.md",
            ROOT / "reports/unified_direct_v5/protocol_amendments.md",
            ROOT / "reports/unified_direct_v5/protocol_freeze.json",
        ]
    missing = [str(path) for path in paths if not path.exists()]
    if missing:
        raise FileNotFoundError("missing provenance file(s): " + ", ".join(missing))
    return {rel(path): sha256(path) for path in paths}


def specified_q1(rows: list[dict[str, str]], strategy: str) -> None:
    folder = OUT / "specified_days/result1"
    folder.mkdir(parents=True, exist_ok=True)
    by_slot = {int(row["slot_id"]): row for row in rows}
    daily_grid = sum(number(row, "grid_kwh") for row in rows)
    daily_cost = sum(number(row, "cost_yuan") for row in rows)
    table1 = []
    for slot in (61, 73, 85, 97, 109, 121):
        row = by_slot[slot]
        table1.append({"record_type": "interval", "strategy_id": strategy, "interval": period((slot - 1) * 10, slot * 10), "grid_kwh": number(row, "grid_kwh"), "daily_grid_kwh": None, "daily_cost_yuan": None})
    table1.append({"record_type": "daily_total", "strategy_id": strategy, "interval": "全天", "grid_kwh": None, "daily_grid_kwh": daily_grid, "daily_cost_yuan": daily_cost})
    table2 = []
    for block in range(6):
        chunk = rows[block * 24:(block + 1) * 24]
        table2.append({"record_type": "four_hour_block", "strategy_id": strategy, "time": period(block * 240, (block + 1) * 240), "charge_kwh": sum(number(row, "charge_kwh") for row in chunk), "discharge_kwh": sum(number(row, "discharge_kwh") for row in chunk), "energy_kwh": None})
    table2 += [
        {"record_type": "storage", "strategy_id": strategy, "time": "00:00", "charge_kwh": None, "discharge_kwh": None, "energy_kwh": number(rows[0], "energy_start_kwh")},
        {"record_type": "storage", "strategy_id": strategy, "time": "24:00", "charge_kwh": None, "discharge_kwh": None, "energy_kwh": number(rows[-1], "energy_end_kwh")},
    ]
    for stem, title, records in (("table1", "问题1表1", table1), ("table2", "问题1表2", table2)):
        write_csv(folder / f"{stem}.csv", records)
        write_markdown(folder / f"{stem}.md", title, records)


def specified_year(name: str, rows: list[dict[str, str]], strategy: str, q2: bool) -> None:
    folder = OUT / f"specified_days/{name}"
    folder.mkdir(parents=True, exist_ok=True)
    _, by_key = keyed_year(rows)
    for date in SPECIFIED_DATES:
        day = [by_key[(date, slot)] for slot in range(1, 145)]
        plan_key = "grid_plan_kwh" if q2 else "grid_original_kwh"
        original_fee_key = "planned_cost_yuan" if q2 else "original_cost_yuan"
        effective_key = plan_key if q2 else "grid_effective_kwh"
        contract_key = "planned_cost_yuan" if q2 else "contract_cost_yuan"
        table1 = []
        for slot in (61, 73, 85, 97, 109, 121):
            row = by_key[(date, slot)]
            table1.append({
                "record_type": "interval", "date": date, "strategy_id": strategy,
                "interval": period((slot - 1) * 10, slot * 10),
                "original_grid_kwh": number(row, plan_key), "effective_grid_kwh": number(row, effective_key),
                "emergency_kwh": number(row, "emergency_kwh"), "original_cost_yuan": number(row, original_fee_key),
                "increase_cost_yuan": 0.0 if q2 else number(row, "increase_cost_yuan"),
                "decrease_adjustment_yuan": 0.0 if q2 else number(row, "decrease_adjustment_yuan"),
                "contract_cost_yuan": number(row, contract_key), "emergency_cost_yuan": number(row, "emergency_cost_yuan"),
                "total_cost_yuan": number(row, "total_cost_yuan"),
            })
        table1.append({
            "record_type": "daily_total", "date": date, "strategy_id": strategy, "interval": "全天",
            "original_grid_kwh": sum(number(row, plan_key) for row in day),
            "effective_grid_kwh": sum(number(row, effective_key) for row in day),
            "emergency_kwh": sum(number(row, "emergency_kwh") for row in day),
            "original_cost_yuan": sum(number(row, original_fee_key) for row in day),
            "increase_cost_yuan": 0.0 if q2 else sum(number(row, "increase_cost_yuan") for row in day),
            "decrease_adjustment_yuan": 0.0 if q2 else sum(number(row, "decrease_adjustment_yuan") for row in day),
            "contract_cost_yuan": sum(number(row, contract_key) for row in day),
            "emergency_cost_yuan": sum(number(row, "emergency_cost_yuan") for row in day),
            "total_cost_yuan": sum(number(row, "total_cost_yuan") for row in day),
        })
        table2 = []
        for block in range(6):
            chunk = day[block * 24:(block + 1) * 24]
            table2.append({"record_type": "four_hour_block", "date": date, "strategy_id": strategy, "time": period(block * 240, (block + 1) * 240), "charge_kwh": sum(number(row, "charge_actual_kwh") for row in chunk), "discharge_kwh": sum(number(row, "discharge_actual_kwh") for row in chunk), "energy_kwh": None})
        table2 += [
            {"record_type": "storage", "date": date, "strategy_id": strategy, "time": "00:00", "charge_kwh": None, "discharge_kwh": None, "energy_kwh": number(day[0], "energy_start_actual_kwh")},
            {"record_type": "storage", "date": date, "strategy_id": strategy, "time": "24:00", "charge_kwh": None, "discharge_kwh": None, "energy_kwh": number(day[-1], "energy_end_actual_kwh")},
        ]
        events = emergency_events(day)
        table3 = [{"date": date, "strategy_id": strategy, "event_status": "有紧急购电", **event} for event in events]
        if not table3:
            table3 = [{"date": date, "strategy_id": strategy, "event_status": "无紧急购电", "interval": "无", "emergency_kwh": 0.0}]
        for stem, title, records in (("table1", f"{name} {date} 表1", table1), ("table2", f"{name} {date} 表2", table2), ("table3", f"{name} {date} 表3", table3)):
            write_csv(folder / f"{date}_{stem}.csv", records)
            write_markdown(folder / f"{date}_{stem}.md", title, records)


def main() -> int:
    if not SELECTION.exists():
        raise FileNotFoundError(f"selection is not ready: {SELECTION}")
    selection_doc = json.loads(SELECTION.read_text(encoding="utf-8"))
    selected = selection_doc["selected"]
    expected = {"q1", "q2", "q3", "q42", "q43"}
    if set(selected) != expected or selected["q1"] != "baseline" or selected["q2"] != "Q2_D112":
        raise ValueError(f"unexpected selection: {selected}")
    sources = {
        "result1": ("baseline", ROOT / "results/q1/baseline/schedule.csv", "q1"),
        "result2": ("Q2_D112", ROOT / "results/q2_direct_v4/runs/D112/ledger.csv", "q2"),
        "result3": (selected["q3"], ROOT / f"results/unified_direct_v5/runs/{selected['q3']}/ledger.csv", "adjusted"),
        "result4-2": (selected["q42"], ROOT / f"results/unified_direct_v5/runs/{selected['q42']}/ledger.csv", "basic_new"),
        "result4-3": (selected["q43"], ROOT / f"results/unified_direct_v5/runs/{selected['q43']}/ledger.csv", "adjusted"),
    }
    OUT.mkdir(parents=True, exist_ok=True)
    REPORT.mkdir(parents=True, exist_ok=True)
    labels = interval_labels()
    payload: list[dict] = []
    mapping: list[dict] = []
    costs: list[dict] = []
    policy_mapping: dict[str, dict] = {}
    for name, (strategy, ledger, kind) in sources.items():
        template = TEMPLATES / f"{name}.xlsx"
        if not ledger.exists():
            raise FileNotFoundError(f"selected ledger is not ready: {ledger}")
        rows = read_csv(ledger)
        q1 = kind == "q1"
        q2 = kind == "q2"
        adjusted_output = kind == "adjusted"
        item = {"name": name, "strategy": strategy, "source": str(ledger), "template": str(template), "labels": labels, "sheets": {}}
        with closing(openpyxl.load_workbook(template, read_only=True, data_only=False)) as workbook:
            for sheet_name in ("计划购电量", "调整购电量"):
                if sheet_name not in workbook.sheetnames:
                    continue
                sheet = workbook[sheet_name]
                for slot, label in enumerate(labels, 1):
                    cell = sheet.cell(slot + 1, 1) if q1 else sheet.cell(1, slot + 1)
                    mapping.append({"file": f"{name}.xlsx", "sheet": sheet_name, "template_cell": cell.coordinate, "original_label": cell.value, "review_label": label, "source_slot_id": slot, "start_minute": (slot - 1) * 10, "end_minute": slot * 10, "operation": "relabel_copy_only_no_data_shift"})
        if q1:
            by_start = {int(row["start_minute"]): row for row in rows}
            if len(rows) != 144 or set(by_start) != set(range(0, 1440, 10)):
                raise ValueError("Q1 interval coverage mismatch")
            rows = [by_start[minute] for minute in range(0, 1440, 10)]
            item["sheets"]["计划购电量"] = [[label, number(row, "grid_kwh")] for label, row in zip(labels, rows)]
            item["sheets"]["充放电量"] = [[period(block * 240, (block + 1) * 240), sum(number(row, "charge_kwh") for row in rows[block * 24:(block + 1) * 24]), sum(number(row, "discharge_kwh") for row in rows[block * 24:(block + 1) * 24]), "00:00" if block == 0 else "24:00" if block == 1 else None, number(rows[0], "energy_start_kwh") if block == 0 else number(rows[-1], "energy_end_kwh") if block == 1 else None] for block in range(6)]
            initial_energy = number(rows[0], "energy_start_kwh")
            final_energy = number(rows[-1], "energy_end_kwh")
            specified_q1(rows, strategy)
            date_range = "single undated representative day"
            omitted_energy = omitted_cost = 0.0
        else:
            dates, by_key = keyed_year(rows)
            plan_key = "grid_plan_kwh" if q2 else "grid_original_kwh"
            effective_key = plan_key if q2 else "grid_effective_kwh"
            original_fee = "planned_cost_yuan" if q2 else "original_cost_yuan"
            contract_fee = "planned_cost_yuan" if q2 else "contract_cost_yuan"
            plans, adjusted, batteries, events = [], [], [], []
            omitted_energy = omitted_cost = 0.0
            for date in dates:
                day = [by_key[(date, slot)] for slot in range(1, 145)]
                plans.append([date] + [number(row, plan_key) for row in day] + [None, sum(number(row, original_fee) for row in day)])
                if adjusted_output:
                    adjusted.append([date] + [number(row, effective_key) for row in day] + [None, sum(number(row, contract_fee) for row in day)])
                for block in range(6):
                    chunk = day[block * 24:(block + 1) * 24]
                    batteries.append([date, period(block * 240, (block + 1) * 240), sum(number(row, "charge_actual_kwh") for row in chunk), sum(number(row, "discharge_actual_kwh") for row in chunk), "00:00" if block == 0 else "24:00" if block == 1 else None, number(day[0], "energy_start_actual_kwh") if block == 0 else number(day[-1], "energy_end_actual_kwh") if block == 1 else None])
                for event in emergency_events(day):
                    events.append([date, event["interval"], event["emergency_kwh"]])
                for row in day:
                    emergency = number(row, "emergency_kwh")
                    if 0.0 < emergency <= 1e-7:
                        omitted_energy += emergency
                        omitted_cost += number(row, "emergency_cost_yuan")
            item["sheets"] = {"计划购电量": plans}
            if adjusted_output:
                item["sheets"]["调整购电量"] = adjusted
            item["sheets"].update({"充放电量": batteries, "紧急购电量": events})
            initial_energy = number(rows[0], "energy_start_actual_kwh")
            final_energy = number(rows[-1], "energy_end_actual_kwh")
            specified_year(name, rows, strategy, q2)
            date_range = "2025-02-01 through 2025-12-31"
        fees = component_sums(rows, q1, q2)
        costs.append({"file": f"{name}.xlsx", "strategy_id": strategy, "days": 1 if q1 else 334, **fees})
        provenance = hashes_for(name, ledger, template, strategy)
        policy_mapping[f"{name}.xlsx"] = {
            "strategy_id": strategy,
            "source_ledger": rel(ledger),
            "date_range": date_range,
            "initial_energy_kwh": initial_energy,
            "final_energy_kwh": final_energy,
            "fee_components_yuan": fees,
            "event_export_threshold_kwh": 1e-7,
            "omitted_tiny_emergency_kwh": omitted_energy,
            "omitted_tiny_emergency_cost_yuan": omitted_cost,
            "sha256": provenance,
            "working_assumptions": {
                "interval_labels": "00:00-00:10 through 23:50-24:00; source values are keyed by physical interval and are never shifted",
                "precision": "numeric values are stored at source precision and displayed to six decimals; displayed values can differ from stored values by rounding",
                "battery": "AC-bus-side charge/discharge; 0.9 efficiency in each direction; energy bounds 1200 to 10800 kWh",
                "billing": "full ledger bill includes every emergency tail; only the event listing omits intervals at or below 1e-7 kWh",
                "interpretation": "protocol.md and protocol_amendments.md contain the frozen working interpretation and export clarification",
            },
        }
        payload.append(item)
    write_csv(OUT / "template_time_mapping.csv", mapping)
    write_csv(OUT / "cost_summary.csv", costs)
    (OUT / "workbook_payload.json").write_text(json.dumps(payload, ensure_ascii=False, allow_nan=False), encoding="utf-8")
    policy_mapping["selection"] = {"path": rel(SELECTION), "sha256": sha256(SELECTION), "selected": selected, "rule": selection_doc.get("rule")}
    (OUT / "policy_mapping.json").write_text(json.dumps(policy_mapping, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    (OUT / "README.md").write_text(
        "# Result workbook exports\n\n"
        "The five workbooks retain source numeric precision. Excel displays energy and fee values to six decimals, so copied display text can lose digits even though the stored cell values remain unrounded. "
        "`policy_mapping.json` records the selected strategies, sources, fees, assumptions, and hashes.\n",
        encoding="utf-8",
    )
    print(json.dumps({"status": "prepared", "selected": selected, "costs": costs}, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

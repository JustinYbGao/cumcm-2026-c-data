#!/usr/bin/env python3
"""Independent readback of saved v5 result workbooks.

The verifier does not import the workbook preparer or builder and does not read
their payload.  Expected bills are reconstructed from the original Appendix 1
and Appendix 4 price workbooks.
"""
from __future__ import annotations

import csv
import datetime as dt
import hashlib
import json
import math
import sys
from collections import defaultdict
from contextlib import closing
from pathlib import Path

import openpyxl


ROOT = Path(__file__).resolve().parents[2]
REPO = ROOT.parent
OUT = ROOT / "outputs/unified_direct_v5"
REPORT = ROOT / "reports/unified_direct_v5"
ATTACHMENTS = REPO / "CUMCM2026Problems/C题/附件"
TEMPLATES = ATTACHMENTS / "附件5"
ENERGY_TOL = 1e-6
COST_TOL = 1e-5
EVENT_THRESHOLD = 1e-7
SPECIFIED_DATES = ("2025-03-20", "2025-06-21", "2025-09-23", "2025-12-21")
BAD_ERRORS = {"#NULL!", "#DIV/0!", "#VALUE!", "#REF!", "#NAME?", "#NUM!", "#N/A", "#SPILL!", "#CALC!"}


def read_csv(path: Path) -> list[dict[str, str]]:
    with path.open(encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1 << 20), b""):
            digest.update(block)
    return digest.hexdigest()


def number(row: dict[str, str], key: str) -> float:
    return float(row[key])


def clock(minute: int) -> str:
    return "24:00" if minute == 1440 else f"{minute // 60:02d}:{minute % 60:02d}"


def period(start: int, end: int) -> str:
    return f"{clock(start)}-{clock(end)}"


def labels() -> list[str]:
    return [period(i * 10, (i + 1) * 10) for i in range(144)]


def date_text(value) -> str:
    if isinstance(value, (dt.datetime, dt.date)):
        return value.strftime("%Y-%m-%d")
    return str(value)[:10]


class Audit:
    def __init__(self) -> None:
        self.checks = 0
        self.numeric = 0
        self.failures = 0
        self.errors: list[dict] = []
        self.max_abs_error = {"kWh": 0.0, "yuan": 0.0, "kW": 0.0}
        self.counts = defaultdict(int)

    def check(self, condition: bool, code: str, location: str, expected=None, actual=None) -> None:
        self.checks += 1
        self.counts[code] += 1
        if not condition:
            self.failures += 1
        if not condition and len(self.errors) < 500:
            record = {"code": code, "location": location}
            if expected is not None:
                record["expected"] = expected
            if actual is not None:
                record["actual"] = actual
            self.errors.append(record)

    def close(self, actual, expected, location: str, unit: str = "kWh") -> None:
        self.numeric += 1
        tolerance = COST_TOL if unit == "yuan" else ENERGY_TOL
        try:
            a, e = float(actual), float(expected)
            error = abs(a - e)
            finite = math.isfinite(a) and math.isfinite(e)
        except (TypeError, ValueError):
            error, finite = math.inf, False
        if finite:
            self.max_abs_error[unit] = max(self.max_abs_error[unit], error)
        self.check(finite and error <= tolerance, "numeric_match", location, expected, actual)


def raw_prices() -> tuple[dict[int, float], dict[tuple[str, int], float], dict[str, str]]:
    p1 = ATTACHMENTS / "附件1.xlsx"
    p4 = ATTACHMENTS / "附件4.xlsx"
    with closing(openpyxl.load_workbook(p1, read_only=True, data_only=True)) as wb:
        ws = wb[wb.sheetnames[0]]
        fixed = {slot: float(values[1]) for slot, values in enumerate(ws.iter_rows(min_row=2, values_only=True), 1)}
    with closing(openpyxl.load_workbook(p4, read_only=True, data_only=True)) as wb:
        ws = wb[wb.sheetnames[0]]
        variable = {}
        for values in ws.iter_rows(min_row=2, values_only=True):
            date = date_text(values[0])
            for slot, value in enumerate(values[1:145], 1):
                variable[(date, slot)] = float(value)
    return fixed, variable, {str(p1.relative_to(REPO)): sha256(p1), str(p4.relative_to(REPO)): sha256(p4)}


def raw_actual_energy() -> tuple[dict[tuple[str, int], tuple[float, float]], str]:
    path = ATTACHMENTS / "附件2.xlsx"
    result: dict[tuple[str, int], tuple[float, float]] = {}
    with closing(openpyxl.load_workbook(path, read_only=True, data_only=True)) as wb:
        load, pv = wb["小区负载"], wb["光伏发电实际功率"]
        load_rows = load.iter_rows(min_row=2, values_only=True)
        pv_rows = pv.iter_rows(min_row=2, values_only=True)
        for load_values, pv_values in zip(load_rows, pv_rows, strict=True):
            date = date_text(load_values[0])
            if date_text(pv_values[0]) != date:
                raise ValueError(f"Appendix 2 date mismatch: {date} vs {date_text(pv_values[0])}")
            for slot, (load_kw, pv_kw) in enumerate(zip(load_values[1:145], pv_values[1:145], strict=True), 1):
                result[(date, slot)] = (float(load_kw) / 6.0, float(pv_kw) / 6.0)
    return result, sha256(path)


def keyed_ledger(rows: list[dict[str, str]], audit: Audit, name: str, actual: dict) -> tuple[list[str], dict[tuple[str, int], dict[str, str]]]:
    keyed = {}
    for row in rows:
        date, slot = row["date"], int(row["slot_id"])
        key = (date, slot)
        audit.check(key not in keyed, "unique_ledger_key", f"{name}/ledger/{date}/{slot}")
        keyed[key] = row
        start = dt.datetime.fromisoformat(row["interval_start"])
        end = dt.datetime.fromisoformat(row["interval_end"])
        expected_start = dt.datetime.fromisoformat(date) + dt.timedelta(minutes=(slot - 1) * 10)
        audit.check(start == expected_start, "ledger_start", f"{name}/ledger/{date}/{slot}", expected_start.isoformat(" "), row["interval_start"])
        audit.check(end == expected_start + dt.timedelta(minutes=10), "ledger_end", f"{name}/ledger/{date}/{slot}")
        if key in actual:
            audit.close(number(row, "load_actual_kwh"), actual[key][0], f"{name}/ledger/load/{date}/{slot}")
            audit.close(number(row, "pv_actual_kwh"), actual[key][1], f"{name}/ledger/pv/{date}/{slot}")
    dates = [(dt.date(2025, 2, 1) + dt.timedelta(days=i)).isoformat() for i in range(334)]
    expected = {(date, slot) for date in dates for slot in range(1, 145)}
    audit.check(set(keyed) == expected, "ledger_coverage", f"{name}/ledger", len(expected), len(keyed))
    return dates, keyed


def independent_components(row: dict[str, str], price: float, q2: bool) -> dict[str, float]:
    emergency = 5.0 * price * number(row, "emergency_kwh")
    if q2:
        original = price * number(row, "grid_plan_kwh")
        return {"original": original, "increase": 0.0, "decrease": 0.0, "contract": original, "emergency": emergency, "total": original + emergency}
    q0 = number(row, "grid_original_kwh")
    effective = number(row, "grid_effective_kwh")
    original = price * q0
    increase = 1.5 * price * max(effective - q0, 0.0)
    decrease = -0.5 * price * max(q0 - effective, 0.0)
    contract = original + increase + decrease
    return {"original": original, "increase": increase, "decrease": decrease, "contract": contract, "emergency": emergency, "total": contract + emergency}


def validate_ledger_math(audit: Audit, name: str, rows: list[dict[str, str]], q2: bool, price_for) -> dict[str, float]:
    totals = defaultdict(float)
    for row in rows:
        date, slot = row["date"], int(row["slot_id"])
        price = price_for(date, slot)
        parts = independent_components(row, price, q2)
        ledger_fields = {"original": "planned_cost_yuan" if q2 else "original_cost_yuan", "contract": "planned_cost_yuan" if q2 else "contract_cost_yuan", "emergency": "emergency_cost_yuan", "total": "total_cost_yuan"}
        if not q2:
            ledger_fields.update({"increase": "increase_cost_yuan", "decrease": "decrease_adjustment_yuan"})
        for component, field in ledger_fields.items():
            audit.close(number(row, field), parts[component], f"{name}/ledger/{field}/{date}/{slot}", "yuan")
        if "price_yuan_per_kwh" in row:
            audit.close(number(row, "price_yuan_per_kwh"), price, f"{name}/ledger/source_price/{date}/{slot}", "yuan")
        for key, value in parts.items():
            totals[key] += value
        start = number(row, "energy_start_actual_kwh")
        charge = number(row, "charge_actual_kwh")
        discharge = number(row, "discharge_actual_kwh")
        end = start + 0.9 * charge - discharge / 0.9
        audit.close(number(row, "energy_end_actual_kwh"), end, f"{name}/ledger/SOC/{date}/{slot}")
        audit.check(1200.0 - ENERGY_TOL <= end <= 10800.0 + ENERGY_TOL, "soc_bounds", f"{name}/ledger/SOC/{date}/{slot}")
        audit.check(-ENERGY_TOL <= charge <= 5000 / 6 + ENERGY_TOL, "charge_bounds", f"{name}/ledger/charge/{date}/{slot}")
        audit.check(-ENERGY_TOL <= discharge <= 5000 / 6 + ENERGY_TOL, "discharge_bounds", f"{name}/ledger/discharge/{date}/{slot}")
        audit.check(not (charge > ENERGY_TOL and discharge > ENERGY_TOL), "no_simultaneous_charge_discharge", f"{name}/ledger/{date}/{slot}")
    return dict(totals)


def validate_plan(audit: Audit, name: str, sheet, cached, dates: list[str], keyed: dict, adjusted: bool, q2: bool, price_for) -> None:
    audit.check((sheet.max_row, sheet.max_column) == (335, 147), "plan_shape", f"{name}/{sheet.title}", [335, 147], [sheet.max_row, sheet.max_column])
    audit.check([sheet.cell(1, col).value for col in range(2, 146)] == labels(), "plan_labels", f"{name}/{sheet.title}")
    audit.check([date_text(sheet.cell(row, 1).value) for row in range(2, 336)] == dates, "plan_dates", f"{name}/{sheet.title}")
    key = "grid_effective_kwh" if adjusted else ("grid_plan_kwh" if q2 else "grid_original_kwh")
    cost_component = "contract" if adjusted else "original"
    for row_number, date in enumerate(dates, 2):
        day = [keyed[(date, slot)] for slot in range(1, 145)]
        for column, ledger_row in enumerate(day, 2):
            cell = sheet.cell(row_number, column)
            audit.check(cell.value is not None and math.isfinite(float(cell.value)), "grid_cell_finite", f"{name}/{sheet.title}!{cell.coordinate}")
            audit.close(cell.value, number(ledger_row, key), f"{name}/{sheet.title}!{cell.coordinate}")
            audit.check(cell.number_format == "0.000000", "six_decimal_format", f"{name}/{sheet.title}!{cell.coordinate}", "0.000000", cell.number_format)
        formula = f"=SUM(B{row_number}:EO{row_number})"
        audit.check(sheet.cell(row_number, 146).value == formula, "daily_amount_formula", f"{name}/{sheet.title}!EP{row_number}", formula, sheet.cell(row_number, 146).value)
        audit.close(cached.cell(row_number, 146).value, sum(number(row, key) for row in day), f"{name}/{sheet.title}!EP{row_number}[cached]")
        expected_cost = sum(independent_components(row, price_for(date, int(row["slot_id"])), q2)[cost_component] for row in day)
        audit.close(sheet.cell(row_number, 147).value, expected_cost, f"{name}/{sheet.title}!EQ{row_number}", "yuan")
        audit.close(cached.cell(row_number, 147).value, expected_cost, f"{name}/{sheet.title}!EQ{row_number}[cached]", "yuan")
        audit.check(sheet.cell(row_number, 147).number_format == "0.000000", "six_decimal_format", f"{name}/{sheet.title}!EQ{row_number}")


def validate_battery(audit: Audit, name: str, sheet, dates: list[str], keyed: dict) -> None:
    audit.check((sheet.max_row, sheet.max_column) == (2005, 6), "battery_shape", f"{name}/充放电量", [2005, 6], [sheet.max_row, sheet.max_column])
    for day_index, date in enumerate(dates):
        day = [keyed[(date, slot)] for slot in range(1, 145)]
        for block in range(6):
            row_number = 2 + day_index * 6 + block
            chunk = day[block * 24:(block + 1) * 24]
            audit.check(date_text(sheet.cell(row_number, 1).value) == date, "battery_date", f"{name}/充放电量!A{row_number}")
            audit.check(sheet.cell(row_number, 2).value == period(block * 240, (block + 1) * 240), "battery_period", f"{name}/充放电量!B{row_number}")
            audit.close(sheet.cell(row_number, 3).value, sum(number(row, "charge_actual_kwh") for row in chunk), f"{name}/充放电量!C{row_number}")
            audit.close(sheet.cell(row_number, 4).value, sum(number(row, "discharge_actual_kwh") for row in chunk), f"{name}/充放电量!D{row_number}")
            expected_mark = "00:00" if block == 0 else "24:00" if block == 1 else None
            expected_energy = number(day[0], "energy_start_actual_kwh") if block == 0 else number(day[-1], "energy_end_actual_kwh") if block == 1 else None
            audit.check(sheet.cell(row_number, 5).value == expected_mark, "battery_soc_label", f"{name}/充放电量!E{row_number}", expected_mark, sheet.cell(row_number, 5).value)
            if expected_energy is None:
                audit.check(sheet.cell(row_number, 6).value is None, "intentional_sparse_soc", f"{name}/充放电量!F{row_number}")
            else:
                audit.close(sheet.cell(row_number, 6).value, expected_energy, f"{name}/充放电量!F{row_number}")


def event_rows(day: list[dict[str, str]]) -> list[tuple[str, float]]:
    result = []
    i = 0
    while i < 144:
        if number(day[i], "emergency_kwh") <= EVENT_THRESHOLD:
            i += 1
            continue
        start, amount = i, 0.0
        while i < 144 and number(day[i], "emergency_kwh") > EVENT_THRESHOLD:
            amount += number(day[i], "emergency_kwh")
            i += 1
        result.append((period(start * 10, i * 10), amount))
    return result


def validate_events(audit: Audit, name: str, sheet, dates: list[str], keyed: dict) -> tuple[float, float, int]:
    expected = []
    omitted_energy = omitted_cost = 0.0
    for date in dates:
        day = [keyed[(date, slot)] for slot in range(1, 145)]
        expected.extend((date, interval, amount) for interval, amount in event_rows(day))
        for row in day:
            energy = number(row, "emergency_kwh")
            if 0.0 < energy <= EVENT_THRESHOLD:
                omitted_energy += energy
                omitted_cost += number(row, "emergency_cost_yuan")
    actual = [(date_text(sheet.cell(row, 1).value), sheet.cell(row, 2).value, sheet.cell(row, 3).value) for row in range(2, sheet.max_row + 1)]
    audit.check(len(actual) == len(expected), "event_count", f"{name}/紧急购电量", len(expected), len(actual))
    for index, expected_row in enumerate(expected[:len(actual)]):
        actual_row = actual[index]
        audit.check(actual_row[:2] == expected_row[:2], "event_key", f"{name}/紧急购电量/{index + 2}", expected_row[:2], actual_row[:2])
        audit.close(actual_row[2], expected_row[2], f"{name}/紧急购电量!C{index + 2}")
    return omitted_energy, omitted_cost, len(expected)


def validate_q1(audit: Audit, workbook, rows: list[dict[str, str]], fixed: dict[int, float]) -> dict[str, float]:
    by_slot = {int(row["slot_id"]): row for row in rows}
    audit.check(len(by_slot) == 144, "q1_coverage", "result1/ledger")
    plan = workbook["计划购电量"]
    audit.check((plan.max_row, plan.max_column) == (145, 2), "q1_plan_shape", "result1/计划购电量")
    audit.check([plan.cell(row, 1).value for row in range(2, 146)] == labels(), "q1_labels", "result1/计划购电量")
    total_grid = total_cost = 0.0
    for slot in range(1, 145):
        expected = number(by_slot[slot], "grid_kwh")
        audit.close(plan.cell(slot + 1, 2).value, expected, f"result1/计划购电量!B{slot + 1}")
        audit.check(plan.cell(slot + 1, 2).number_format == "0.000000", "six_decimal_format", f"result1/计划购电量!B{slot + 1}")
        total_grid += expected
        total_cost += fixed[slot] * expected
        audit.close(number(by_slot[slot], "cost_yuan"), fixed[slot] * expected, f"result1/ledger/cost/{slot}", "yuan")
    battery = workbook["充放电量"]
    for block in range(6):
        chunk = [by_slot[slot] for slot in range(block * 24 + 1, block * 24 + 25)]
        row_number = block + 2
        audit.check(battery.cell(row_number, 1).value == period(block * 240, (block + 1) * 240), "q1_battery_period", f"result1/充放电量!A{row_number}")
        audit.close(battery.cell(row_number, 2).value, sum(number(row, "charge_kwh") for row in chunk), f"result1/充放电量!B{row_number}")
        audit.close(battery.cell(row_number, 3).value, sum(number(row, "discharge_kwh") for row in chunk), f"result1/充放电量!C{row_number}")
        mark = "00:00" if block == 0 else "24:00" if block == 1 else None
        audit.check(battery.cell(row_number, 4).value == mark, "q1_soc_label", f"result1/充放电量!D{row_number}")
    audit.close(battery.cell(2, 5).value, number(by_slot[1], "energy_start_kwh"), "result1/充放电量!E2")
    audit.close(battery.cell(3, 5).value, number(by_slot[144], "energy_end_kwh"), "result1/充放电量!E3")
    audit.close(total_cost, 35126.94858928963, "result1/reference_total", "yuan")
    return {"original": total_cost, "increase": 0.0, "decrease": 0.0, "contract": total_cost, "emergency": 0.0, "total": total_cost, "grid": total_grid}


def parse_markdown(path: Path) -> list[list[str]]:
    lines = [line for line in path.read_text(encoding="utf-8").splitlines() if line.startswith("|")]
    if len(lines) < 2:
        return []
    return [[cell.strip() for cell in line.strip().strip("|").split("|")] for line in lines[2:]]


def markdown_cell(csv_value: str) -> str:
    if csv_value == "":
        return ""
    try:
        return f"{float(csv_value):.6f}"
    except ValueError:
        return csv_value.replace("|", "\\|")


def validate_specified_files(audit: Audit, name: str, strategy: str, q1_rows=None, dates=None, keyed=None, q2=False, price_for=None) -> None:
    folder = OUT / f"specified_days/{name}"
    pairs = [(folder / "table1.csv", folder / "table1.md"), (folder / "table2.csv", folder / "table2.md")] if name == "result1" else [(folder / f"{date}_{table}.csv", folder / f"{date}_{table}.md") for date in SPECIFIED_DATES for table in ("table1", "table2", "table3")]
    for csv_path, md_path in pairs:
        audit.check(csv_path.exists(), "specified_csv_exists", str(csv_path))
        audit.check(md_path.exists(), "specified_md_exists", str(md_path))
        if not csv_path.exists() or not md_path.exists():
            continue
        csv_rows = read_csv(csv_path)
        markdown_rows = parse_markdown(md_path)
        audit.check(len(markdown_rows) == len(csv_rows), "specified_markdown_rows", str(md_path), len(csv_rows), len(markdown_rows))
        for index, row in enumerate(csv_rows[:len(markdown_rows)]):
            audit.check(markdown_rows[index][list(row).index("strategy_id")] == strategy, "specified_strategy_md", f"{md_path}/{index}")
            audit.check(row["strategy_id"] == strategy, "specified_strategy_csv", f"{csv_path}/{index}")
            audit.check(markdown_rows[index] == [markdown_cell(row[field]) for field in row], "specified_markdown_values", f"{md_path}/{index}")
    if name == "result1":
        by_slot = {int(row["slot_id"]): row for row in q1_rows}
        t1 = read_csv(folder / "table1.csv")
        audit.check(len(t1) == 7, "specified_q1_table1_rows", str(folder / "table1.csv"))
        for row, slot in zip(t1[:6], (61, 73, 85, 97, 109, 121)):
            audit.check(row["interval"] == period((slot - 1) * 10, slot * 10), "specified_q1_interval", row["interval"])
            audit.close(row["grid_kwh"], number(by_slot[slot], "grid_kwh"), f"specified/result1/table1/{slot}")
        audit.close(t1[-1]["daily_grid_kwh"], sum(number(row, "grid_kwh") for row in q1_rows), "specified/result1/daily_grid")
        audit.close(t1[-1]["daily_cost_yuan"], sum(fixed for fixed in [number(row, "cost_yuan") for row in q1_rows]), "specified/result1/daily_cost", "yuan")
        t2 = read_csv(folder / "table2.csv")
        audit.check(len(t2) == 8, "specified_q1_table2_rows", str(folder / "table2.csv"))
        return
    for date in SPECIFIED_DATES:
        day = [keyed[(date, slot)] for slot in range(1, 145)]
        plan_key = "grid_plan_kwh" if q2 else "grid_original_kwh"
        effective_key = plan_key if q2 else "grid_effective_kwh"
        t1 = read_csv(folder / f"{date}_table1.csv")
        audit.check(len(t1) == 7, "specified_table1_rows", f"{name}/{date}")
        for row, slot in zip(t1[:6], (61, 73, 85, 97, 109, 121)):
            source = keyed[(date, slot)]
            audit.check(row["interval"] == period((slot - 1) * 10, slot * 10), "specified_interval", f"{name}/{date}/{slot}")
            audit.close(row["original_grid_kwh"], number(source, plan_key), f"{name}/{date}/table1/original/{slot}")
            audit.close(row["effective_grid_kwh"], number(source, effective_key), f"{name}/{date}/table1/effective/{slot}")
            audit.close(row["emergency_kwh"], number(source, "emergency_kwh"), f"{name}/{date}/table1/emergency/{slot}")
            component = independent_components(source, price_for(date, slot), q2)
            for csv_key, part_key in (("original_cost_yuan", "original"), ("increase_cost_yuan", "increase"), ("decrease_adjustment_yuan", "decrease"), ("contract_cost_yuan", "contract"), ("emergency_cost_yuan", "emergency"), ("total_cost_yuan", "total")):
                audit.close(row[csv_key], component[part_key], f"{name}/{date}/table1/{csv_key}/{slot}", "yuan")
        daily = t1[-1]
        audit.close(daily["original_grid_kwh"], sum(number(row, plan_key) for row in day), f"{name}/{date}/table1/daily_original")
        audit.close(daily["effective_grid_kwh"], sum(number(row, effective_key) for row in day), f"{name}/{date}/table1/daily_effective")
        audit.close(daily["emergency_kwh"], sum(number(row, "emergency_kwh") for row in day), f"{name}/{date}/table1/daily_emergency")
        parts = [independent_components(row, price_for(date, int(row["slot_id"])), q2) for row in day]
        for csv_key, component in (("original_cost_yuan", "original"), ("increase_cost_yuan", "increase"), ("decrease_adjustment_yuan", "decrease"), ("contract_cost_yuan", "contract"), ("emergency_cost_yuan", "emergency"), ("total_cost_yuan", "total")):
            audit.close(daily[csv_key], sum(part[component] for part in parts), f"{name}/{date}/table1/{csv_key}", "yuan")
        t2 = read_csv(folder / f"{date}_table2.csv")
        audit.check(len(t2) == 8, "specified_table2_rows", f"{name}/{date}")
        for block in range(6):
            chunk = day[block * 24:(block + 1) * 24]
            audit.close(t2[block]["charge_kwh"], sum(number(row, "charge_actual_kwh") for row in chunk), f"{name}/{date}/table2/charge/{block}")
            audit.close(t2[block]["discharge_kwh"], sum(number(row, "discharge_actual_kwh") for row in chunk), f"{name}/{date}/table2/discharge/{block}")
        audit.close(t2[6]["energy_kwh"], number(day[0], "energy_start_actual_kwh"), f"{name}/{date}/table2/start")
        audit.close(t2[7]["energy_kwh"], number(day[-1], "energy_end_actual_kwh"), f"{name}/{date}/table2/end")
        t3 = read_csv(folder / f"{date}_table3.csv")
        expected_events = event_rows(day)
        if expected_events:
            audit.check(len(t3) == len(expected_events), "specified_table3_rows", f"{name}/{date}")
            for row, expected in zip(t3, expected_events):
                audit.check(row["event_status"] == "有紧急购电" and row["interval"] == expected[0], "specified_table3_event", f"{name}/{date}/{expected[0]}")
                audit.close(row["emergency_kwh"], expected[1], f"{name}/{date}/table3/{expected[0]}")
        else:
            audit.check(len(t3) == 1 and t3[0]["event_status"] == "无紧急购电" and t3[0]["interval"] == "无", "specified_table3_zero", f"{name}/{date}")
            audit.close(t3[0]["emergency_kwh"], 0.0, f"{name}/{date}/table3/zero")


def validate_mapping(audit: Audit) -> None:
    rows = read_csv(OUT / "template_time_mapping.csv")
    audit.check(len(rows) == 1008, "mapping_count", "template_time_mapping.csv", 1008, len(rows))
    expected = set()
    for filename, sheets in {
        "result1.xlsx": ("计划购电量",), "result2.xlsx": ("计划购电量",),
        "result3.xlsx": ("计划购电量", "调整购电量"), "result4-2.xlsx": ("计划购电量",),
        "result4-3.xlsx": ("计划购电量", "调整购电量"),
    }.items():
        with closing(openpyxl.load_workbook(TEMPLATES / filename, read_only=False, data_only=False)) as wb:
            for sheet_name in sheets:
                for slot in range(1, 145):
                    expected.add((filename, sheet_name, slot))
                    match = [row for row in rows if row["file"] == filename and row["sheet"] == sheet_name and int(row["source_slot_id"]) == slot]
                    audit.check(len(match) == 1, "mapping_unique", f"{filename}/{sheet_name}/{slot}")
                    if not match:
                        continue
                    row = match[0]
                    cell = f"A{slot + 1}" if filename == "result1.xlsx" else f"{openpyxl.utils.get_column_letter(slot + 1)}1"
                    audit.check(row["template_cell"] == cell, "mapping_cell", f"{filename}/{sheet_name}/{slot}", cell, row["template_cell"])
                    audit.check(row["original_label"] == str(wb[sheet_name][cell].value), "mapping_original", f"{filename}/{sheet_name}/{slot}")
                    audit.check(row["review_label"] == labels()[slot - 1], "mapping_review", f"{filename}/{sheet_name}/{slot}")
                    audit.check((int(row["start_minute"]), int(row["end_minute"])) == ((slot - 1) * 10, slot * 10), "mapping_minutes", f"{filename}/{sheet_name}/{slot}")
    audit.check({(row["file"], row["sheet"], int(row["source_slot_id"])) for row in rows} == expected, "mapping_coverage", "template_time_mapping.csv")


def validate_metadata(audit: Audit, mapping: dict, filename: str, strategy: str, ledger: Path, totals: dict, omitted: tuple[float, float]) -> None:
    record = mapping.get(filename)
    audit.check(record is not None, "metadata_present", filename)
    if not record:
        return
    audit.check(record["strategy_id"] == strategy, "metadata_strategy", filename, strategy, record["strategy_id"])
    audit.check((ROOT / record["source_ledger"]).resolve() == ledger.resolve(), "metadata_ledger", filename)
    fee = record["fee_components_yuan"]
    for metadata_key, component in (("original_cost_yuan", "original"), ("increase_cost_yuan", "increase"), ("decrease_adjustment_yuan", "decrease"), ("contract_cost_yuan", "contract"), ("emergency_cost_yuan", "emergency"), ("total_cost_yuan", "total")):
        audit.close(fee[metadata_key], totals[component], f"metadata/{filename}/{metadata_key}", "yuan")
    audit.close(record["omitted_tiny_emergency_kwh"], omitted[0], f"metadata/{filename}/omitted_energy")
    audit.close(record["omitted_tiny_emergency_cost_yuan"], omitted[1], f"metadata/{filename}/omitted_cost", "yuan")
    audit.check(record["event_export_threshold_kwh"] == EVENT_THRESHOLD, "metadata_event_threshold", filename)
    audit.check("displayed to six decimals" in record["working_assumptions"]["precision"], "metadata_precision_statement", filename)
    for path_text, recorded_hash in record["sha256"].items():
        path = (REPO if path_text.startswith("CUMCM2026Problems/") else ROOT) / path_text
        audit.check(path.exists(), "metadata_hash_path", path_text)
        if path.exists():
            audit.check(sha256(path) == recorded_hash, "metadata_hash", path_text, recorded_hash, sha256(path))


def scan_workbook(audit: Audit, name: str, workbook, cached) -> None:
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value in BAD_ERRORS:
                    audit.check(False, "formula_error", f"{name}/{sheet.title}!{cell.coordinate}", actual=cell.value)
                cached_value = cached[sheet.title][cell.coordinate].value
                if isinstance(cached_value, str) and cached_value in BAD_ERRORS:
                    audit.check(False, "cached_formula_error", f"{name}/{sheet.title}!{cell.coordinate}", actual=cached_value)


def main() -> int:
    audit = Audit()
    fixed, variable, price_hashes = raw_prices()
    actual, actual_hash = raw_actual_energy()
    selection = json.loads((ROOT / "results/unified_direct_v5/selection.json").read_text(encoding="utf-8"))["selected"]
    policy_mapping = json.loads((OUT / "policy_mapping.json").read_text(encoding="utf-8"))
    selection_record = policy_mapping.get("selection", {})
    audit.check(selection_record.get("selected") == selection, "metadata_selection", "policy_mapping.json")
    selection_path = ROOT / selection_record.get("path", "missing")
    audit.check(selection_path.exists(), "metadata_selection_path", str(selection_path))
    if selection_path.exists():
        audit.check(selection_record.get("sha256") == sha256(selection_path), "metadata_selection_hash", str(selection_path))
    config = {
        "result1": ("baseline", ROOT / "results/q1/baseline/schedule.csv", "q1", False, False),
        "result2": ("Q2_D112", ROOT / "results/q2_direct_v4/runs/D112/ledger.csv", "fixed", True, False),
        "result3": (selection["q3"], ROOT / f"results/unified_direct_v5/runs/{selection['q3']}/ledger.csv", "fixed", False, True),
        "result4-2": (selection["q42"], ROOT / f"results/unified_direct_v5/runs/{selection['q42']}/ledger.csv", "variable", False, False),
        "result4-3": (selection["q43"], ROOT / f"results/unified_direct_v5/runs/{selection['q43']}/ledger.csv", "variable", False, True),
    }
    workbook_results = {}
    for name, (strategy, ledger, price_kind, q2, adjusted_output) in config.items():
        filename = f"{name}.xlsx"
        output = OUT / filename
        rows = read_csv(ledger)
        workbook = openpyxl.load_workbook(output, data_only=False)
        cached = openpyxl.load_workbook(output, data_only=True)
        scan_workbook(audit, name, workbook, cached)
        expected_sheets = ["计划购电量", "充放电量"] if name == "result1" else ["计划购电量", "调整购电量", "充放电量", "紧急购电量"] if adjusted_output else ["计划购电量", "充放电量", "紧急购电量"]
        audit.check(workbook.sheetnames == expected_sheets, "sheet_order", name, expected_sheets, workbook.sheetnames)
        if name == "result1":
            totals = validate_q1(audit, workbook, rows, fixed)
            omitted = (0.0, 0.0)
            validate_specified_files(audit, name, strategy, q1_rows=rows)
            event_count = 0
        else:
            dates, keyed = keyed_ledger(rows, audit, name, actual)
            price_for = (lambda date, slot: fixed[slot]) if price_kind == "fixed" else (lambda date, slot: variable[(date, slot)])
            totals = validate_ledger_math(audit, name, rows, q2, price_for)
            validate_plan(audit, name, workbook["计划购电量"], cached["计划购电量"], dates, keyed, False, q2, price_for)
            if adjusted_output:
                validate_plan(audit, name, workbook["调整购电量"], cached["调整购电量"], dates, keyed, True, q2, price_for)
            validate_battery(audit, name, workbook["充放电量"], dates, keyed)
            omitted_energy, omitted_cost, event_count = validate_events(audit, name, workbook["紧急购电量"], dates, keyed)
            omitted = (omitted_energy, omitted_cost)
            validate_specified_files(audit, name, strategy, dates=dates, keyed=keyed, q2=q2, price_for=price_for)
            if name == "result2":
                audit.close(totals["total"], 13913892.481868185, "result2/D112_reference_total", "yuan")
        validate_metadata(audit, policy_mapping, filename, strategy, ledger, totals, omitted)
        workbook_results[filename] = {
            "strategy_id": strategy,
            "xlsx_sha256": sha256(output),
            "source_ledger_sha256": sha256(ledger),
            "event_rows": event_count,
            "omitted_tiny_emergency_kwh": omitted[0],
            "omitted_tiny_emergency_cost_yuan": omitted[1],
            "fee_components_yuan": totals,
            "ledger_intervals": len(rows),
            "dates": 1 if name == "result1" else 334,
            "grid_cells_read": 144 if name == "result1" else 48096 * (2 if adjusted_output else 1),
        }
        workbook.close()
        cached.close()
    validate_mapping(audit)
    summary_rows = {row["file"]: row for row in read_csv(OUT / "cost_summary.csv")}
    for filename, result in workbook_results.items():
        audit.check(filename in summary_rows, "cost_summary_row", filename)
        if filename in summary_rows:
            row = summary_rows[filename]
            for csv_key, component in (("original_cost_yuan", "original"), ("increase_cost_yuan", "increase"), ("decrease_adjustment_yuan", "decrease"), ("contract_cost_yuan", "contract"), ("emergency_cost_yuan", "emergency"), ("total_cost_yuan", "total")):
                audit.close(row[csv_key], result["fee_components_yuan"][component], f"cost_summary/{filename}/{csv_key}", "yuan")
    report = {
        "passed": audit.failures == 0,
        "status": "pass" if audit.failures == 0 else "fail",
        "units": {"energy": "kWh", "power": "kW", "cost": "yuan", "price": "yuan/kWh"},
        "tolerances": {"energy_abs_kwh": ENERGY_TOL, "cost_abs_yuan": COST_TOL, "event_listing_threshold_kwh": EVENT_THRESHOLD},
        "checks": audit.checks,
        "numeric_comparisons": audit.numeric,
        "error_count": audit.failures,
        "max_abs_error": audit.max_abs_error,
        "check_counts": dict(audit.counts),
        "coverage": {
            "workbooks": len(workbook_results),
            "ledger_intervals": sum(result["ledger_intervals"] for result in workbook_results.values()),
            "workbook_grid_cells": sum(result["grid_cells_read"] for result in workbook_results.values()),
            "year_dates": 334 * 4,
            "template_time_mappings": 1008,
            "specified_day_csv_files": 2 + 4 * 4 * 3,
            "specified_day_markdown_files": 2 + 4 * 4 * 3,
        },
        "workbooks": workbook_results,
        "raw_source_hashes": {**price_hashes, "CUMCM2026Problems/C题/附件/附件2.xlsx": actual_hash},
        "validator_sha256": sha256(Path(__file__)),
        "independence": "Read actual saved XLSX files with openpyxl and reconstructed costs from original Appendix 1/4 price workbooks. The Artifact Tool and workbook_payload.json were not imported or read.",
        "errors": audit.errors,
    }
    REPORT.mkdir(parents=True, exist_ok=True)
    (REPORT / "workbook_readback.json").write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    lines = [
        "# Workbook independent readback", "", f"- Status: **{report['status'].upper()}**", f"- Checks: {audit.checks}",
        f"- Numeric comparisons: {audit.numeric}", f"- Errors: {audit.failures}",
        f"- Maximum absolute energy error: {audit.max_abs_error['kWh']:.12g} kWh",
        f"- Maximum absolute cost error: {audit.max_abs_error['yuan']:.12g} yuan", "",
        "Saved workbooks were read independently with openpyxl. Bills were rebuilt from the original Appendix 1 and Appendix 4 prices using the applicable original contract, F_A adjustment, and 5p emergency formulas.", "",
        "| Workbook | Strategy | Events | Omitted tiny emergency (kWh) | Total cost (yuan) |", "| --- | --- | ---: | ---: | ---: |",
    ]
    for filename, result in workbook_results.items():
        lines.append(f"| {filename} | {result['strategy_id']} | {result['event_rows']} | {result['omitted_tiny_emergency_kwh']:.12g} | {result['fee_components_yuan']['total']:.6f} |")
    if audit.errors:
        lines += ["", "## Errors", ""] + [f"- `{error['code']}` at `{error['location']}`: expected `{error.get('expected')}`, actual `{error.get('actual')}`" for error in audit.errors[:100]]
    (REPORT / "workbook_readback.md").write_text("\n".join(lines) + "\n", encoding="utf-8")
    print(json.dumps({"passed": report["passed"], "status": report["status"], "checks": audit.checks, "numeric_comparisons": audit.numeric, "error_count": audit.failures, "max_abs_error": audit.max_abs_error, "coverage": report["coverage"]}, ensure_ascii=False))
    return 0 if audit.failures == 0 else 1


if __name__ == "__main__":
    sys.exit(main())

"""Read-only inventory of official inputs, templates and the friend's notebooks."""
from pathlib import Path
from collections import Counter
import hashlib
import json
import math
import openpyxl
from pypdf import PdfReader

WORK = Path(__file__).resolve().parents[1]
ROOT = WORK.parent
SOURCE = ROOT / 'CUMCM2026Problems/C题/附件'


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def main():
    report = {}
    manifest = {}
    for path in sorted(SOURCE.rglob('*.xlsx')):
        if path.name.startswith('~$'):
            continue  # Excel lock files are not attachments.
        rel = str(path.relative_to(ROOT))
        manifest[rel] = sha(path)
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        sheets = {}
        for ws in wb:
            rows = list(ws.values)
            cells = [(r, c, v) for r, row in enumerate(rows, 1)
                     for c, v in enumerate(row, 1) if v is not None and v != '']
            nums = [v for _, _, v in cells if isinstance(v, (int, float))]
            sheets[ws.title] = {
                'dimensions': [ws.max_row, ws.max_column],
                'types': dict(Counter(type(v).__name__ for _, _, v in cells)),
                'first_rows': rows[:3], 'last_row': rows[-1],
                'empty_or_blank_cells': sum(v is None or v == '' for row in rows for v in row),
                'numeric_min': min(nums) if nums else None,
                'numeric_max': max(nums) if nums else None,
                'nonfinite': sum(not math.isfinite(v) for v in nums),
            }
        report[rel] = sheets
        wb.close()
    (WORK / 'data/interim/source_inventory.json').write_text(
        json.dumps(report, ensure_ascii=False, indent=2, default=str))
    (WORK / 'data/interim/source_manifest.json').write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2))
    pdf = ROOT / 'CUMCM2026Problems/C题/C题.pdf'
    (WORK / 'data/interim/problem_text.txt').write_text('\n\n'.join(
        f'PAGE {i+1}\n{p.extract_text()}' for i, p in enumerate(PdfReader(pdf).pages)))
    notebook_summary = {}
    for path in sorted((ROOT / '111').glob('*.ipynb')):
        nb = json.loads(path.read_text())
        extracted, outputs = [], []
        for i, cell in enumerate(nb['cells']):
            if cell['cell_type'] == 'code':
                extracted.append(f'\n# CELL {i}, execution_count={cell.get("execution_count")}\n'
                                 + ''.join(cell['source']))
                for out in cell.get('outputs', []):
                    outputs.append({'cell': i, 'type': out['output_type'],
                                    'text': out.get('text', out.get('data', {}).get('text/plain', [])),
                                    'error': out.get('evalue')})
        (WORK / 'data/interim' / (path.stem + '.py')).write_text('\n'.join(extracted))
        notebook_summary[path.name] = outputs
    (WORK / 'data/interim/notebook_outputs.json').write_text(
        json.dumps(notebook_summary, ensure_ascii=False, indent=2))
    for path, sheets in report.items():
        print(path)
        for name, s in sheets.items():
            print(name, s['dimensions'], 'types=', s['types'])
            print('first=', str(s['first_rows'])[:1100])
            print('last=', str(s['last_row'])[:260])
    print('Inspection saved; original files opened read-only.')


if __name__ == '__main__':
    main()

"""Descriptive plots and evidence reports, without fitted predictors or optimization."""
from pathlib import Path
import json
import os
import platform
import importlib.metadata
import numpy as np
import pandas as pd
import openpyxl
from openpyxl.utils import get_column_letter

WORK = Path(__file__).resolve().parents[1]
ROOT = WORK.parent
os.environ['MPLCONFIGDIR'] = str(WORK/'data/interim/matplotlib')
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
import matplotlib.dates as mdates


def template_report():
    records, summary = [], ['# 附件5只读检查', '', '本轮未填写或修改模板。题面附录2规定电量与储电量单位均为kWh。', '',
        '| 文件/工作表 | 维度 | 原模板非空结构 |', '|---|---|---|']
    for p in sorted((ROOT/'CUMCM2026Problems/C题/附件/附件5').glob('*.xlsx')):
        wb = openpyxl.load_workbook(p,read_only=True,data_only=True)
        for ws in wb:
            nonempty = []
            for row in ws:
                for cell in row:
                    if cell.value is not None:
                        records.append({'file':p.name,'sheet':ws.title,'cell':cell.coordinate,
                                        'value':str(cell.value),'python_type':type(cell.value).__name__})
                        nonempty.append(f'{cell.coordinate}={cell.value}')
            if ws.title in ['计划购电量','调整购电量']:
                if p.name=='result1.xlsx':
                    detail = f'A2={ws["A2"].value}；A145={ws["A145"].value}；144条购电时段'
                else:
                    detail = f'B1={ws["B1"].value}；EO1={ws["EO1"].value}；EP1={ws["EP1"].value}；EQ1={ws["EQ1"].value}；日期{ws["A2"].value}至{ws["A335"].value}'
            else:
                detail = '；'.join(nonempty[:10])+'；…（完整坐标见template_cells.csv）'
            summary.append(f'| {p.name}/{ws.title} | {ws.max_row}×{ws.max_column} | {detail} |')
        wb.close()
    pd.DataFrame(records).to_csv(WORK/'reports/template_cells.csv',index=False)
    summary += ['', '## 真实问题与下游要求', '',
        '- 计划/调整购电模板从00:10–00:20开始，末段结束于次日00:10，相对00:00–24:00口径偏移10分钟。result1首行A2和末行A145可直接复核；年度模板首末时段在B1、EO1。年度EO1原文为0:00-0:10+1，起点还缺少+1标记，与result1末行写法不一致。应确认含义，不能通过强行平移原始数值掩盖。',
        '- 年度计划/调整表含334个业务日（2025-02-01至2025-12-31），另有全天购电量/购电费汇总列，不能将147列都当时间列。',
        '- 充放电量、紧急购电量工作表包含少量日期及省略号，是示意结构，题面要求2025-02-01至12-31完整结果；后续填写须扩展，不能只填模板已有示例行。',
        '- 论文指定四天为2025-03-20、06-21、09-23、12-21，不等于只需模拟四天。Q1另需全天电量与费用、六段4小时充放电及0/24点储电量。',
        '- result3/result4-3须区分0点计划与调整后电量；当前数据处理中只保留预报版本，不产生购电承诺或SOC数值。',
        '- 模板时间偏移的最终填表映射需要队伍核实题意/澄清说明后决定，本轮不修改模板。']
    (WORK/'reports/template_review.md').write_text('\n'.join(summary)+'\n')


def main():
    a = pd.read_csv(WORK/'data/processed/actual_10min.csv',float_precision='round_trip',parse_dates=['interval_start','interval_end'])
    q = pd.read_csv(WORK/'data/processed/q1_day.csv',float_precision='round_trip')
    h = pd.read_csv(WORK/'data/processed/pv_forecast_hourly.csv',float_precision='round_trip',parse_dates=['issue_time','target_time'])
    d = pd.read_csv(WORK/'data/processed/pv_forecast_10min.csv',float_precision='round_trip')
    daily = a.groupby('date')[['load_actual_kwh','pv_actual_kwh','net_load_actual_kwh']].sum()
    daily.to_csv(WORK/'reports/daily_energy.csv')
    calendar = daily.reset_index()
    dates = pd.to_datetime(calendar.date)
    calendar['month'] = dates.dt.month
    calendar['weekday'] = dates.dt.dayofweek
    calendar_rows = []
    for grouping in ['month','weekday']:
        for value,g in calendar.groupby(grouping):
            calendar_rows.append({'grouping':grouping,'value':value,'n_days':len(g),
                'mean_daily_load_kwh':g.load_actual_kwh.mean(),
                'mean_daily_pv_kwh':g.pv_actual_kwh.mean(),
                'mean_daily_net_load_kwh':g.net_load_actual_kwh.mean()})
    pd.DataFrame(calendar_rows).to_csv(WORK/'reports/actual_calendar_summary.csv',index=False)
    stats = []
    for name,frame,cols in [('q1_day',q,['load_kw','pv_forecast_kw','price_yuan_per_kwh','net_load_forecast_kwh']),
                            ('actual_10min',a,['load_actual_kw','pv_actual_kw','fixed_price_yuan_per_kwh','actual_price_yuan_per_kwh','net_load_actual_kwh']),
                            ('pv_forecast_hourly',h,['pv_forecast_kw'])]:
        for col in cols:
            x = frame[col]
            stats.append({'table':name,'field':col,'count':len(x),'missing':int(x.isna().sum()),
                          'nonfinite':int((~np.isfinite(x)).sum()),'zero':int(x.eq(0).sum()),
                          'negative':int(x.lt(0).sum()),'min':x.min(),'max':x.max()})
    pd.DataFrame(stats).to_csv(WORK/'reports/numeric_summary.csv',index=False)
    jumps = []
    for col,prefix in [('load_actual_kw','load'),('pv_actual_kw','pv'),('actual_price_yuan_per_kwh','price')]:
        diff = a[col].diff()
        for idx in diff.abs().nlargest(10).index:
            jumps.append({'field':col,'interval_start':a.loc[idx,'interval_start'],
                          'previous_value':a.loc[idx-1,col],'current_value':a.loc[idx,col],
                          'change':diff.loc[idx], 'source_row':a.loc[idx,f'{prefix}_source_row'],
                          'source_column':a.loc[idx,f'{prefix}_source_column'],
                          'action':'保留；仅按全样本相邻绝对变化列出前10条，不判为错误'})
    pd.DataFrame(jumps).to_csv(WORK/'reports/largest_jumps.csv',index=False)
    plt.rcParams.update({'font.size':10,'axes.spines.top':False,'axes.spines.right':False,'figure.dpi':140})
    fig,axes = plt.subplots(2,2,figsize=(12,7),layout='constrained')
    for ax,day in zip(axes.flat,['2025-03-20','2025-06-21','2025-09-23','2025-12-21']):
        subset = a.loc[a.date==day]
        ax.plot((subset.slot_id-1)/6,subset.load_actual_kw,label='Load',color='#2864ad')
        ax.plot((subset.slot_id-1)/6,subset.pv_actual_kw,label='PV',color='#e59326')
        ax.set(title=day,ylabel='Average power (kW)')
        ax.set_xlim(0,24)
        ax.set_xticks([0,6,12,18,24],['00:00','06:00','12:00','18:00','24:00'])
        ax.legend(frameon=False)
    fig.suptitle('Observed daily profiles — interval-average assumption')
    fig.savefig(WORK/'reports/figures/representative_days.png')
    plt.close(fig)
    fig,ax = plt.subplots(figsize=(12,4),layout='constrained')
    for col,label,color in [('load_actual_kwh','Load','#2864ad'),('pv_actual_kwh','PV','#e59326')]:
        ax.plot(pd.to_datetime(daily.index),daily[col],label=label,color=color,linewidth=1)
    ax.set(title='2025 observed daily energy — sum of power / 6',ylabel='Energy (kWh/day)',xlabel='Business date')
    ax.legend(frameon=False)
    fig.savefig(WORK/'reports/figures/daily_energy.png')
    plt.close(fig)
    fig,axes = plt.subplots(1,2,figsize=(12,4),layout='constrained')
    subset = a.loc[a.date=='2025-03-20']
    axes[0].plot((subset.slot_id-1)/6,subset.actual_price_yuan_per_kwh,label='Actual',color='#2864ad')
    axes[0].plot((subset.slot_id-1)/6,subset.fixed_price_yuan_per_kwh,label='Fixed',color='#e59326')
    axes[0].set_xlim(0,24)
    axes[0].set_xticks([0,6,12,18,24],['00:00','06:00','12:00','18:00','24:00'])
    axes[0].set(title='2025-03-20 price comparison',ylabel='Price (yuan/kWh)')
    axes[0].legend(frameon=False)
    axes[1].boxplot([g.actual_price_yuan_per_kwh for _,g in a.groupby(a.interval_start.dt.month)],showfliers=False)
    axes[1].set(title='Monthly actual prices (outlier dots hidden)',xlabel='Month',ylabel='Price (yuan/kWh)')
    fig.savefig(WORK/'reports/figures/prices.png')
    plt.close(fig)
    template_report()
    validations = json.loads((WORK/'reports/validation_all.json').read_text())
    diagnostic_validations = json.loads((WORK/'reports/validation_forecast_diagnostics.json').read_text())
    lines = ['# 数据质量报告','',f'基础数据重新读回后执行 {len(validations)} 项检查，通过 {sum(v["passed"] for v in validations)} 项。新增预报诊断执行 {len(diagnostic_validations)} 项独立检查，通过 {sum(v["passed"] for v in diagnostic_validations)} 项。明细见 validation_all.md与validation_forecast_diagnostics.json；失败项不能解释为通过。','',
             '| 表 | 实际行数 | 唯一键 |','|---|---:|---|',
             f'| q1_day | {len(q)} | slot_id |','| fixed_price | 144 | slot_id |',
             f'| actual_10min | {len(a)} | interval_start（也可date+slot_id） |',
             f'| pv_forecast_hourly | {len(h)} | issue_time+target_time |',
             f'| pv_forecast_10min | {len(d)} | issue_time+interval_start |','',
             '## 实际发现','',f'- 实际PV零值 {a.pv_actual_kw.eq(0).sum()} 个，有符号净负载为负 {a.net_load_actual_kwh.lt(0).sum()} 个；全部保留。',
             f'- 实际负载范围 {a.load_actual_kw.min()}—{a.load_actual_kw.max()} kW；实际PV范围 {a.pv_actual_kw.min()}—{a.pv_actual_kw.max()} kW；实际价格范围 {a.actual_price_yuan_per_kwh.min()}—{a.actual_price_yuan_per_kwh.max()} 元/kWh。',
             '- 四份输入的数值区域未发现缺失、非数值或非有限值；未发现负功率/负价格。原表合法零值没有填补。数值统计见 numeric_summary.csv，首轮原字段类型见 source_inventory.json。',
             '- 附件3有1095个空字符串日期，按原顺序向下填充，原标签和逐单元格记录都保留；不是1095条预报缺失。',
             f'- 小时预报保留了 {(h.target_time>pd.Timestamp("2026-01-01")).sum()} 条超出实际数据终点的目标，最大目标为 {h.target_time.max()}；未将其实际值填零。',
             '- 第一份预报2025-01-01 00:00缺少可用的0点端点，因此派生表缺00:00–01:00的6个区间。其余1459份各144区间，总210234行。此缺口不影响2—12月的规定输出范围。',
             '- 附件5全部计划/调整表的时段标签相对采用口径偏移10分钟；详见 template_review.md，未修改原模板。',
             '', '## 尖峰与图的解释', '',
             'largest_jumps.csv逐字段列出相邻变化绝对值最大的10条及原始行列。这是描述性排名，没有采用删除阈值，没有证据将它们判为数据错误。三个探索图仅描述观测曲线、日总电量与价格分布；箱线图仅隐藏离群点图标，数据未删除。没有计算或声称预测准确率、节费率。',
             '', '## 补充完成的数据诊断', '',
             'pv_forecast_hourly_evaluation.csv保留35040条版本—目标记录，其中35003条小时电量可比较；不完整首小时与跨年缺实际值分别标记，未填零。pv_forecast_error_summary.csv提供偏差、MAE和RMSE，严格版本比较采用8742个相同目标小时。详见pv_forecast_diagnostics.md。',
             'actual_calendar_summary.csv按月份和星期汇总观测日总电量（星期0=周一，6=周日）。这些全样本统计只作描述，不输入早期决策或训练特征。largest_jumps.csv保留原有逐值跳变诊断，未清除真实尖峰。',
             '', '## 采用假设与未完成项', '',
             '区间结束标签、区间平均功率、实际区间结束即完整可用、整点预报点值积分及旧版端点均为待队伍确认的处理假设，详见processing_notes.md。未来实际价格未标为提前可知。未训练最终模型、未求解调度、未填写策略结果。']
    (WORK/'reports/quality_report.md').write_text('\n'.join(lines)+'\n')
    versions = {'python':platform.python_version()}
    for package in ['numpy','pandas','openpyxl','pypdf','matplotlib','ipython']:
        versions[package] = importlib.metadata.version(package)
    (WORK/'logs/runtime_versions.json').write_text(json.dumps(versions,indent=2))
    print('\n'.join(lines[:15]))
    print('Saved three descriptive figures, template audit, numeric/jump summaries and quality report.')


if __name__=='__main__':
    main()

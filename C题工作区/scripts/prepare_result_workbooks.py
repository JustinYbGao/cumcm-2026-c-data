"""Prepare keyed, traceable template payloads; never writes an Excel workbook."""
import csv
import hashlib
import json
from collections import defaultdict
from datetime import datetime, timedelta
from pathlib import Path

import openpyxl

WORK = Path(__file__).resolve().parents[1]
OUT = WORK / 'outputs/revision_v1'
INTERIM = WORK / 'data/interim/revision_v1'
SOURCES = {
    'result1': 'results/q1/baseline/schedule.csv',
    'result2': 'results/q2/selected/ledger.csv',
    'result3': 'results/robustness/fixed_w28/ledger.csv',
    'result4-2': 'results/q4/q42_ols/ledger.csv',
    'result4-3': 'results/robustness/variable_w28/ledger.csv',
}


def clock(minute):
    return f'{minute // 60:02d}:{minute % 60:02d}'


def period(start, end):
    return f'{clock(start)}-{clock(end)}'


def number(row, key):
    return float(row[key])


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    INTERIM.mkdir(parents=True, exist_ok=True)
    payload, mapping, costs, hashes = [], [], [], {}
    for name, source in SOURCES.items():
        source_path = WORK / source
        template = WORK.parent / f'CUMCM2026Problems/C题/附件/附件5/{name}.xlsx'
        for path in (source_path, template):
            hashes[str(path.relative_to(WORK.parent))] = hashlib.sha256(path.read_bytes()).hexdigest()
        with source_path.open() as f:
            rows = list(csv.DictReader(f))
        wb = openpyxl.load_workbook(template, read_only=True)
        item = {'name': name, 'source': source, 'template': str(template), 'sheets': {}}
        labels = [period(t, t + 10) for t in range(0, 1440, 10)]
        for sheet_name in ('计划购电量', '调整购电量'):
            if sheet_name not in wb.sheetnames:
                continue
            planned = wb[sheet_name]
            for slot, label in enumerate(labels, 1):
                cell = planned.cell(slot + 1, 1) if name == 'result1' else planned.cell(1, slot + 1)
                mapping.append({'file': name + '.xlsx', 'sheet': sheet_name, 'template_cell': cell.coordinate,
                                'original_label': cell.value, 'review_label': label,
                                'source_slot_id': slot, 'start_minute': (slot - 1) * 10,
                                'end_minute': slot * 10, 'operation': 'relabel_copy_only_no_data_shift'})
        if name == 'result1':
            keyed = {int(r['start_minute']): r for r in rows}
            assert len(rows) == len(keyed) == 144 and set(keyed) == set(range(0, 1440, 10))
            rows = [keyed[t] for t in range(0, 1440, 10)]
            item['sheets']['计划购电量'] = [[label, number(r, 'grid_kwh')] for label, r in zip(labels, rows)]
            item['sheets']['充放电量'] = [
                [period(b * 240, (b + 1) * 240),
                 sum(number(r, 'charge_kwh') for r in rows[b * 24:(b + 1) * 24]),
                 sum(number(r, 'discharge_kwh') for r in rows[b * 24:(b + 1) * 24]),
                 '00:00' if b == 0 else '24:00' if b == 1 else None,
                 number(rows[0], 'energy_start_kwh') if b == 0 else number(rows[-1], 'energy_end_kwh') if b == 1 else None]
                for b in range(6)]
            costs.append({'file': name + '.xlsx', 'policy': 'baseline', 'days': 1,
                          'original_cost_yuan': sum(number(r, 'cost_yuan') for r in rows),
                          'increase_cost_yuan': 0, 'decrease_adjustment_yuan': 0,
                          'contract_cost_yuan': sum(number(r, 'cost_yuan') for r in rows),
                          'emergency_cost_yuan': 0, 'total_cost_yuan': sum(number(r, 'cost_yuan') for r in rows)})
        else:
            daily = defaultdict(dict)
            for r in rows:
                start = datetime.fromisoformat(r['interval_start'])
                end = datetime.fromisoformat(r['interval_end'])
                slot = start.hour * 6 + start.minute // 10
                assert start.second == 0 and start.minute % 10 == 0 and end - start == timedelta(minutes=10)
                assert r['date'] == start.date().isoformat() and int(r['slot_id']) == slot + 1
                assert slot not in daily[r['date']]
                daily[r['date']][slot] = r
            expected_days = [(datetime(2025, 2, 1) + timedelta(days=i)).date().isoformat() for i in range(334)]
            assert sorted(daily) == expected_days and len(rows) == 48096
            is_q2 = name == 'result2'
            revised = name in ('result3', 'result4-3')
            q0 = 'grid_plan_kwh' if is_q2 else 'grid_original_kwh'
            original_fee = 'planned_cost_yuan' if is_q2 else 'original_cost_yuan'
            contract_fee = 'planned_cost_yuan' if is_q2 else 'contract_cost_yuan'
            plan_rows, adjusted_rows, battery_rows, emergency_rows = [], [], [], []
            for day in expected_days:
                assert set(daily[day]) == set(range(144))
                rs = [daily[day][i] for i in range(144)]
                plan_rows.append([day] + [number(r, q0) for r in rs] + [None, sum(number(r, original_fee) for r in rs)])
                if revised:
                    adjusted_rows.append([day] + [number(r, 'grid_effective_kwh') for r in rs] + [None, sum(number(r, contract_fee) for r in rs)])
                for b in range(6):
                    battery_rows.append([day, period(b * 240, (b + 1) * 240),
                        sum(number(r, 'charge_actual_kwh') for r in rs[b * 24:(b + 1) * 24]),
                        sum(number(r, 'discharge_actual_kwh') for r in rs[b * 24:(b + 1) * 24]),
                        '00:00' if b == 0 else '24:00' if b == 1 else None,
                        number(rs[0], 'energy_start_actual_kwh') if b == 0 else number(rs[-1], 'energy_end_actual_kwh') if b == 1 else None])
                i = 0
                while i < 144:
                    if number(rs[i], 'emergency_kwh') <= 1e-7:
                        i += 1
                        continue
                    start, energy = i, 0.
                    while i < 144 and number(rs[i], 'emergency_kwh') > 1e-7:
                        energy += number(rs[i], 'emergency_kwh')
                        i += 1
                    emergency_rows.append([day, period(start * 10, i * 10), energy])
            item['labels'] = labels
            item['sheets'] = {'计划购电量': plan_rows}
            if revised:
                item['sheets']['调整购电量'] = adjusted_rows
            item['sheets'].update({'充放电量': battery_rows, '紧急购电量': emergency_rows})
            costs.append({'file': name + '.xlsx', 'policy': source.split('/')[-2], 'days': 334,
                          'original_cost_yuan': sum(number(r, original_fee) for r in rows),
                          'increase_cost_yuan': 0 if is_q2 else sum(number(r, 'increase_cost_yuan') for r in rows),
                          'decrease_adjustment_yuan': 0 if is_q2 else sum(number(r, 'decrease_adjustment_yuan') for r in rows),
                          'contract_cost_yuan': sum(number(r, contract_fee) for r in rows),
                          'emergency_cost_yuan': sum(number(r, 'emergency_cost_yuan') for r in rows),
                          'total_cost_yuan': sum(number(r, 'total_cost_yuan') for r in rows)})
        wb.close()
        payload.append(item)
    (INTERIM / 'workbook_payload.json').write_text(json.dumps(payload, ensure_ascii=False, allow_nan=False))
    for filename, records in [('template_time_mapping.csv', mapping), ('cost_summary.csv', costs)]:
        with (OUT / filename).open('w') as f:
            writer = csv.DictWriter(f, fieldnames=records[0].keys())
            writer.writeheader()
            writer.writerows(records)
    (OUT / 'source_hashes.json').write_text(json.dumps(hashes, ensure_ascii=False, indent=2) + '\n')
    print(json.dumps(costs, indent=2))


if __name__ == '__main__':
    main()

"""Regenerate Q1 figures, time-key audit and model/results note from validated saved files."""
import hashlib
import json
import os
from pathlib import Path
import re
import sys

WORK = Path(__file__).resolve().parents[1]
os.environ['TMPDIR'] = str(WORK/'data/interim/q1')
os.environ['MPLCONFIGDIR'] = str(WORK/'data/interim/q1/cache/matplotlib')
sys.dont_write_bytecode = True
import numpy as np
import openpyxl
import pandas as pd


def load_json(path):
    return json.loads(path.read_text())


def clock(minute):
    return f'{int(minute)//60:02d}:{int(minute)%60:02d}'


def period(start,end):
    return f'{clock(start)}–{clock(end)}'


def main():
    root = WORK/'results/q1'
    validation = load_json(root/'validation.json')
    if not validation['passed']:
        raise RuntimeError('Refuse to report unvalidated results')
    p = pd.read_csv(root/'baseline/schedule.csv',float_precision='round_trip')
    sens = pd.read_csv(root/'roundtrip_90/schedule.csv',float_precision='round_trip')
    baseline = load_json(root/'baseline/summary.json')
    other = load_json(root/'roundtrip_90/summary.json')
    no = load_json(root/'no_storage/summary.json')
    status = load_json(root/'baseline/solver_status.json')
    status2 = load_json(root/'roundtrip_90/solver_status.json')
    nonunique = load_json(root/'nonuniqueness.json')
    environment = load_json(WORK/'logs/q1/environment.json')
    cfg = load_json(root/'baseline/config_snapshot.json')
    # Literal label matching only. Never align template and result by row position.
    template = WORK.parent/'CUMCM2026Problems/C题/附件/附件5/result1.xlsx'
    wb = openpyxl.load_workbook(template,read_only=True,data_only=True)
    sheet = wb['计划购电量']
    audit = []
    def minute(label):
        match = re.fullmatch(r'(\d+):(\d+)(\+1)?',label)
        if not match:
            raise ValueError(f'Unexpected template time: {label}')
        return 60*int(match[1])+int(match[2])+(1440 if match[3] else 0)
    for row in range(2,sheet.max_row+1):
        label = sheet.cell(row,1).value
        a,b = [minute(v) for v in label.split('-')]
        hit = p.loc[(p.start_minute==a)&(p.end_minute==b)]
        audit.append({'template_cell':f'A{row}','original_label':label,'literal_start_minute':a,
                      'literal_end_minute':b,'internal_slot_if_literal_match':int(hit.slot_id.item()) if len(hit)==1 else None,
                      'export_status':'not_exported_pending_confirmation'})
    raw = openpyxl.load_workbook(WORK.parent/'CUMCM2026Problems/C题/附件/附件1.xlsx',read_only=True,data_only=True)
    raw_first,raw_last = str(raw['Sheet1']['A2'].value),str(raw['Sheet1']['A145'].value)
    table2_labels = [wb['充放电量'].cell(r,1).value for r in range(2,8)]
    wb.close()
    raw.close()
    pd.DataFrame(audit).to_csv(root/'template_literal_time_audit.csv',index=False)
    evidence = {'raw_attachment1_A2':raw_first,'raw_attachment1_A145':raw_last,
                'template_A2':audit[0]['original_label'],'template_A145':audit[-1]['original_label'],
                'template_table2_blocks':table2_labels,
                'literal_matches':sum(v['internal_slot_if_literal_match'] is not None for v in audit),
                'internal_missing_from_template':'00:00–00:10', 'template_outside_internal_day':'24:00–24:10',
                'decision':'No workbook exported; internal CSV uses explicit start/end keys; no data movement',
                'english_evidence':'Existing English extraction Appendix 2 explicitly says 0:00 to 24:00; does not resolve supplied template labels',
                'initial_state':'Chinese Appendix 1: Jan 1 6000; English: start of operation 6000; Q1 adopts 6000 without inventing a date'}
    (root/'time_mapping_evidence.json').write_text(json.dumps(evidence,ensure_ascii=False,indent=2)+'\n')
    figures = WORK/'reports/figures/modelviz_v3_en'
    figures.mkdir(parents=True,exist_ok=True)
    from plot_results_modelviz import render
    render('q1')
    t1=pd.read_csv(root/'baseline/table1_intervals.csv')
    t2=pd.read_csv(root/'baseline/table2_blocks.csv')
    table1='| 时间段 | 购电量/kWh | 时间段 | 购电量/kWh | 时间段 | 购电量/kWh |\n|---|---:|---|---:|---|---:|\n'
    for offset in [0,3]:
        values=[]
        for _,row in t1.iloc[offset:offset+3].iterrows(): values += [period(row.start_minute,row.end_minute),f'{row.grid_kwh:.6f}']
        table1+='| '+' | '.join(values)+' |\n'
    table1+=f'| 全天购电量 | {baseline["grid_kwh"]:.6f} | 全天购电费/元 | {baseline["cost_yuan"]:.6f} | | |\n'
    table2='| 时间段 | 充电量/kWh | 放电量/kWh | 时间段 | 充电量/kWh | 放电量/kWh |\n|---|---:|---:|---|---:|---:|\n'
    for offset in [0,2,4]:
        values=[]
        for _,row in t2.iloc[offset:offset+2].iterrows(): values += [period(row.block_start,row.block_end),f'{row.charge_kwh:.6f}',f'{row.discharge_kwh:.6f}']
        table2+='| '+' | '.join(values)+' |\n'
    table2+='| 0:00储电量 | 6000.000000 | | 24:00储电量 | 6000.000000 | |\n'
    block_compare='| 四小时区段 | 基准充电/kWh | 对照充电/kWh | 基准放电/kWh | 对照放电/kWh |\n|---|---:|---:|---:|---:|\n'
    st2=pd.read_csv(root/'roundtrip_90/table2_blocks.csv')
    for i,row in t2.iterrows():
        sr=st2.iloc[i]
        block_compare+=f'| {period(row.block_start,row.block_end)} | {row.charge_kwh:.3f} | {sr.charge_kwh:.3f} | {row.discharge_kwh:.3f} | {sr.discharge_kwh:.3f} |\n'
    max_physics=max(v for k,v in validation['scenarios']['baseline']['residuals'].items() if not k.startswith(('summary','interval_cost')))
    savings=no['cost_yuan']-baseline['cost_yuan']
    improvement=baseline['cost_yuan']-other['cost_yuan']
    checks=sum(len(v['checks']) for v in validation['scenarios'].values())+len(validation['cross_checks'])
    report=fr'''# 问题1：确定性储能调度模型与内部结果

> 状态：按“原始时间为区间终点、功率为区间平均值”的工作假设完成内部求解和验收。模板时间映射尚待确认，未生成正式 result1.xlsx，不是正式提交版。问题2—4未执行。

## 1. 当前阶段、材料与执行范围

当前阶段为“正在求解”；数据清洗与复核已完成，本轮不改 processed 文件。继承《C题-统一建模实施方案（含公式）》第4—5节，使用 BZD 工作流、模型字典 #913 与 #797、模型求解检查 skill。完整中文PDF已重新提取并目视核对第1—2页；英文既有提取与附件5原工作簿用于辅助核对。已有完整题意、统一方案、数据字典及验证材料；缺少的是模板口径确认和完整论文，而非本轮计算输入。

本轮依次完成：核对模型与口径 → MILP与LP求解 → 保存CSV后独立重算 → 无储能与效率对照 → 图表、结果说明及BZD局部审查。模型适配补充见 [q1_model_fit.md](q1_model_fit.md)，局部检查见 [q1_bzd_solution_check.md](q1_bzd_solution_check.md)。未训练预测器，也没有开展全年优化。

## 2. 数据、来源与必须公开的假设

输入 `data/processed/q1_day.csv` 有144行，唯一键为 slot_id；用 start_minute/end_minute 计算区间。价格读取 price_yuan_per_kwh，负荷电量读取 load_kwh，光伏读取 pv_forecast_kwh。原 load_kw、pv_forecast_kw 和源行保留。总负荷 {p.load_kwh.sum():.6f} kWh，总给定PV {p.pv_forecast_kwh.sum():.6f} kWh；所有功率按1/6小时转换，价格不转换。没有缺失填补、平移、重采样、伪造来源日期或改写已复核数据。

| 项目 | 本轮采用值/解释 | 依据与边界 |
|---|---|---|
| 时间 | 00:10原始标签对应[00:00,00:10)，共144段 | 原表00:10至0:00+1；终点标签仍是工作假设 |
| 功率 | 各十分钟平均功率 | 采用已复核数据口径；非点功率积分 |
| 额定容量/运行边界 | 12000 / 1200—10800 kWh | 中文附录1；优化使用运行边界 |
| 功率边界 | 交流母线侧充、放电各5000 kW | 边界所在侧是工作假设；区间上限833.333333… kWh |
| 初始与末态 | 固定6000 kWh，日末相等 | 中文附录1只明确2025-01-01初值；英文称运行初始。对无日期Q1显式采用6000，不让优化器自由选择初值 |
| 基准效率 | 充电0.9，放电0.9，往返0.81 | 题面90%表述不唯一，非官方定论 |
| 效率对照 | 双向均sqrt(0.9)，往返0.9 | 仅改变效率，其余输入和边界不变 |
| 弃光 | 0≤w≤给定PV，无售电、无未供负荷 | 本轮将统一方案的富余处置量收紧为PV弃光，不将电池损耗记作弃光 |
| 预测 | 附件1给定PV直接作为确定输入 | 未检验真实运行预测误差；附件3首小时缺口不涉及问题1 |

时间证据另存 `results/q1/time_mapping_evidence.json` 与 `template_literal_time_audit.csv`。原模板A2为 `{audit[0]['original_label']}`，A145为 `{audit[-1]['original_label']}`，字面覆盖00:10—次日00:10；内部覆盖00:00—24:00。按起止时间键只能找到143个相同区间：内部00:00—00:10无模板位置，模板24:00—24:10无当天内部解。英文附件说明明确0:00—24:00，但没有解释原模板偏移。因此不按行号填表、不循环移动数值，也不默认改模板标签。模板确认后，应单独实施导出并从最终工作簿再次验收。

## 3. 模型与单位

令 t=0,…,143，E₀,…,E₁₄₄ 为145个状态。给定区间负荷 Lₜ、光伏 PVₜ（kWh）和价格 pₜ（元/kWh）。变量 gₜ 为购电量，cₜ 为交流母线输入充电量，dₜ 为电池向母线输出量，wₜ 为弃光量，单位均为kWh；Eₜ为电池内部储能（kWh），zₜ∈{{0,1}}，1允许充电、0允许放电，空闲时不限。

目标及全部约束为：

$$\min J=\sum_{{t=0}}^{{143}}p_tg_t.$$

$$g_t+PV_t-w_t+d_t=L_t+c_t,$$

$$E_{{t+1}}=E_t+\eta_c c_t-d_t/\eta_d,$$

$$1200\le E_t\le10800\quad(t=0,\ldots,144),\qquad E_0=E_{{144}}=6000,$$

$$0\le c_t\le (5000/6)z_t,\qquad 0\le d_t\le (5000/6)(1-z_t),$$

$$g_t\ge0,\quad 0\le w_t\le PV_t,\quad z_t\in\{{0,1\}}.$$

充电损耗为(1−ηc)c，放电损耗为(1/ηd−1)d，均在状态递推中计入，不与弃光混合。等式满足负荷，不允许未供电；PV富余可充电或弃光。没有额外售电、购电功率上限或电池折旧成本，因为本问未给这些参数。区间内准静态平均功率假设不能用于保证秒级动态安全。

模型共 {status['variables']} 个变量（其中144个二元变量）、{status['constraints']} 条线性约束及变量上下界。功率产生紧的互斥上界，不使用任意巨大M。LP只将z改为[0,1]，其他式子保持相同，作为MILP的下界和交叉核验；交付方案来自MILP，不以LP直接替代物理调度。

## 4. 求解、最优性与结果

使用 Python {sys.version.split()[0]}、HiGHS {environment['highs']}，单线程、随机种子0、时间限120秒，相对间隙1e-9、绝对间隙1e-7元，原始/对偶/整数可行容差1e-8。HiGHS通过预求解、分支定界及割平面求解MILP，LP日志显示使用对偶单纯形。版本、日志、模型文本、原始解和配置快照全部落盘；无随机采样算法。

基准目标为 **{baseline['cost_yuan']:.8f}元**，购电 **{baseline['grid_kwh']:.8f} kWh**。求解状态为 {status['status']}，MILP下界 {status['lower_bound_yuan']:.8f}元，LP下界 {status['lp_lower_bound_yuan']:.8f}元，相对最优性间隙 {status['mip_gap']:.1f}，分支节点{status['nodes']}。本次基准MILP求解器计时 {status['runtime_seconds']:.4f}秒（不含数据读取、建模、图表和验证），不外推为全年实际耗时。在上述输入与模型假设、数值容差内，已获得全局最优目标；这不证明未确认的题意口径正确，也不代表真实供需下的全年最优。

表1按题面六个十分钟区间，以真实起止键提取，未按模板行号移动。

{table1}
表2按完整四小时窗口累计交流侧充放电量。同一四小时内可以先充后放，互斥限制作用于各十分钟区间。

{table2}
CSV保留完整精度；正文六位小数仅用于展示，重算不得使用显示时舍入后的数。

## 5. 基线比较与节约来源

无储能对照设c=d=0、E恒为6000，g=max(L−PV,0)、w=max(PV−L,0)，同样满足全部物理约束。

| 方案 | 费用/元 | 购电/kWh | 充电/kWh | 放电/kWh | 弃光/kWh |
|---|---:|---:|---:|---:|---:|
| 无储能 | {no['cost_yuan']:.6f} | {no['grid_kwh']:.6f} | 0 | 0 | {no['curtailment_kwh']:.6f} |
| 基准：单向各90% | {baseline['cost_yuan']:.6f} | {baseline['grid_kwh']:.6f} | {baseline['charge_kwh']:.6f} | {baseline['discharge_kwh']:.6f} | {baseline['curtailment_kwh']:.6f} |
| 对照：往返90% | {other['cost_yuan']:.6f} | {other['grid_kwh']:.6f} | {other['charge_kwh']:.6f} | {other['discharge_kwh']:.6f} | {other['curtailment_kwh']:.6f} |

基准节约 **{savings:.6f}元（{savings/no['cost_yuan']*100:.6f}%）**，购电量减少 {no['grid_kwh']-baseline['grid_kwh']:.6f} kWh。全天电池损耗为充电减放电 **{baseline['charge_kwh']-baseline['discharge_kwh']:.6f} kWh**；减少的弃光 {no['curtailment_kwh']:.6f} kWh，扣掉电池损耗，正好解释购电量减少。额外费用节约来自购电时段转移，不能把节费率等同于节电率。

图1展示日内储能及效率对照：基准状态触及1200和10800 kWh边界，但固定首尾6000，未通过耗尽初始储能获得虚假收益。

![图1：内部工作假设下的储能曲线](figures/modelviz_v3_en/battery_states/outputs/chart.png)

图2将给定供需、购电、充放电和电价对齐。基准有{int((p.charge_kwh>1e-6).sum())}个充电区间和{int((p.discharge_kwh>1e-6).sum())}个放电区间；充电功率有{int((abs(p.charge_kwh-5000/6)<1e-6).sum())}段达到5000 kW，放电峰值 {p.discharge_kwh.max()*6:.4f} kW。白天富余PV得到利用，低价窗口充电、高价窗口放电共同降低费用；这描述返回的计划，并非唯一时间规律。

![图2：基准调度与供需、电价](figures/modelviz_v3_en/dispatch/outputs/chart.png)

## 6. 效率敏感性与非唯一性

对照将单向效率由0.9改为sqrt(0.9)≈0.948683298，其余输入、初末态和边界完全一致，并独立求解MILP及LP。对照MILP与LP均为Optimal，间隙为0，下界 {status2['lp_lower_bound_yuan']:.8f}元。相对基准成本减少 **{improvement:.6f}元（{improvement/baseline['cost_yuan']*100:.4f}%）**，购电减少 {baseline['grid_kwh']-other['grid_kwh']:.6f} kWh；相对无储能节费 {(no['cost_yuan']-other['cost_yuan'])/no['cost_yuan']*100:.6f}%。因此采用哪一种90%解释会实质影响结果，不能省略。

{block_compare}
两种效率下最大单区间购电差 {abs(p.grid_kwh-sens.grid_kwh).max():.6f} kWh，最大期末储能差 {abs(p.energy_end_kwh-sens.energy_end_kwh).max():.6f} kWh；`efficiency_sensitivity.json`还记录逐项L1差。只有两个有依据的口径情景，不声称连续参数稳定区间、概率置信区间或全年鲁棒性。

另在基准费用上加1e-7元的容差上限，最小化Σ(t+1)cₜ，获得独立可行替代计划：实际费用差 {nonunique['cost_difference_yuan']:.3g}元，最大计划/状态差 {nonunique['max_plan_difference_kwh']:.6f} kWh。该证据说明数值最优面上存在明显不同安排；没有符号精确的唯一性证明，不将某组动作称为唯一最优规律，也不把两方案的所有逐时差异都归因于效率。

## 7. 独立验证与测试证据

`validate_q1_solution.py`不导入求解器或求解脚本，重新读取源CSV、每个保存的schedule、145个states、summary、table1、table2、配置快照与状态。独立复算供需、递推、容量、交流侧功率、非负、弃光上限、模式整数性、十分钟互斥、连续状态、固定首尾、费用和汇总，以及基线和效率差异。LP检查允许分数模式，但保留功率联立约束并报告是否同时充放电。

本轮六组结果共 **{checks}项检查通过**（含两组LP与一组费用面探测），基准最大供需残差 {validation['scenarios']['baseline']['residuals']['energy_balance']:.3g} kWh，状态递推残差 {validation['scenarios']['baseline']['residuals']['state_recursion']:.3g} kWh；基准物理/模式约束最大残差 {max_physics:.3g}。验收阈值为电量1e-6 kWh、成本1e-6元、整数1e-7；无超限充放电或未供负荷，首尾均6000 kWh。成本由源价格与导出购电量重新相乘得到，与汇总和最优目标一致。

测试包括手算两区间套利、正常计划，以及故意功率超限、状态递推错误、汇总费用错误、时间错移、非有限数值；另有LP无整数统计量的JSON回归测试。8项测试实际通过，初次失败与修复后日志均保存。最初发现HiGHS变量不支持直接除法，已改为乘倒数；LP整数违反指标不可用，存null；这些是实现/输出修复，没有改变物理模型或输入。完整证据见 `logs/q1/tests.log`、`logs/q1/validate.log` 和 `results/q1/validation.json`。

原附件、原模板、中文PDF与全部processed CSV在求解前后SHA-256一致，验证器又按原清单核对当前文件；未改其他成员内容。环境临时文件、安装缓存和图形缓存定向至 `data/interim/q1/`。安装仅增加highspy==1.14.0，网络受限首试日志和成功安装日志均保留。

## 8. 文件接口、复现与验收边界

`baseline/schedule.csv`的g/c/d/w/E列及明确起止键可作为后续确定性调度接口。`states.csv`有145个明确分钟状态；表1六行和表2六行单独落盘，首尾及全天量费在summary中。`*_milp.lp`、`*_lp.lp`及`.sol`保留实际模型与解，solver_status保存下界、gap与计时。`alternate`仅用于说明数值非唯一性，不能替换已指定的baseline方案。**未导出Excel**：时间确认前，完整CSV、题面格式表和图构成内部验收包；正式工作簿导出仍未完成。

在项目根目录执行：

```bash
cd /Users/justingao/Documents/CUMCM
export PYTHONDONTWRITEBYTECODE=1
export TMPDIR="$PWD/C题工作区/data/interim/q1"
C题工作区/.venv/bin/python C题工作区/scripts/solve_q1.py > C题工作区/logs/q1/solve.log 2>&1
C题工作区/.venv/bin/python C题工作区/scripts/validate_q1_solution.py > C题工作区/logs/q1/validate.log 2>&1
C题工作区/.venv/bin/python -m unittest discover -s C题工作区/tests -p test_q1_solution.py > C题工作区/logs/q1/tests.log 2>&1
C题工作区/.venv/bin/python C题工作区/scripts/report_q1.py > C题工作区/logs/q1/report.log 2>&1
```

每条命令退出码须为0；首次迁移安装依赖仍使用README中的项目内虚拟环境命令。`report_q1.py`依据已验证结果重建图与本文；BZD适配和审查报告在模型、口径或数据改变后须重新检查，不自动视为有效。

## 9. 合并待办、下一步与BZD流程进度

1. 正式输出前确认模板首末区间；确认后单独编写明确映射的Excel导出并读回验证。
2. 团队确认效率单向/往返、功率所在侧、区间平均功率、Q1初值采用口径；若改变需重算本轮所有结果。
3. 问题2可进入开发，但需先落实1月因果初始化、跨日SOC和日末条件、历史信息预测及执行规则。附件3第一版缺首小时仅在实际涉及它时制定因果回退，不在本轮补值。
4. 补完整论文时再做跨问误差传递、参考附录、格式和AI使用记录复核，不对未完成论文评分。

| 阶段 | 所需材料 | Skill/执行器 | 状态 | 产物 | 阻塞项 | 下一步 |
|---|---|---|---|---|---|---|
| 完整题意与方案 | 中文PDF、已有分析 | 继承已有BZD成果 | 已完成 | 选题分析目录 | 无本轮计算阻塞 | 保持统一路线 |
| 数据清洗复核 | 原附件、processed | 继承已有数据审计 | 已完成 | 数据字典与复核报告 | 时间语义待最终确认 | 不重复清洗 |
| Q1模型适配 | 方案、144区间 | model-dictionary | 已完成 | q1_model_fit.md、字典记录 | 假设待确认 | 口径变动再审 |
| Q1内部求解 | 附件1与参数 | MILP + LP | 已完成 | results/q1/ | 无内部阻塞 | 保留结果 |
| Q1验证与效率比较 | 保存的结果 | 独立验证器 | 已完成 | validation.json、对照结果 | 真实预测误差不在本轮 | 后问另验 |
| Q1局部章节检查 | 本报告、代码、结果 | model-solution-checker | 已完成 | q1_bzd_solution_check.md | 无完整论文 | 扩稿后复核 |
| Q1正式模板导出 | 已确认时间映射 | 尚未执行 | 未开始 | 无正式result1.xlsx | 时间映射 | 确认后导出再验 |
| Q2 | 历史信息、初始执行方案 | 预测→调度→回放 | 未开始 | 本轮无产物 | 初始化及执行口径 | 可进入开发 |
| Q3—Q4 | Q2、版本预报与计费口径 | 后续MPC等 | 未开始 | 本轮无产物 | 前序工作 | 本轮不展开 |
| 全文与终稿审查 | 完整论文 | BZD后续检查 | 未开始 | 无评分 | 缺完整论文 | 完稿后处理 |

AI协助范围：本轮由Codex读取本地材料与指定skill，核对字典，编写求解、验证、测试与绘图脚本，实际执行计算并起草本说明及局部检查。没有声称人工独立复算、真实运行试验、外部文献检索或全年策略已完成；“独立验证”指与求解代码分离的读回与公式重算，不等于独立人员审计。团队应据实际使用情况将这些记录纳入最终AI使用声明。
'''
    (WORK/'reports/q1_model_and_results.md').write_text(report)
    artifacts = list(root.rglob('*')) + [p for name in ['battery_states', 'dispatch'] for p in (figures/name).rglob('*') if p.is_file()]
    artifacts += [WORK/f for f in ['scripts/solve_q1.py','scripts/validate_q1_solution.py','scripts/report_q1.py',
                                   'tests/test_q1_solution.py','configs/model_baseline.json','requirements.txt','README.md']]
    artifacts += list((WORK/'reports').glob('q1_*.md'))
    manifest = {str(f.relative_to(WORK)):hashlib.sha256(f.read_bytes()).hexdigest() for f in sorted(artifacts) if f.is_file()}
    (WORK/'logs/q1/artifact_hashes.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2)+'\n')
    print('Generated report, two figures and explicit time mapping evidence; no Excel submission export.')


if __name__ == '__main__':
    main()
